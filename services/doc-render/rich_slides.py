"""Shared rich slide geometry; vector PDF and native editable PPTX objects."""

from __future__ import annotations

import io
from typing import Any

from reportlab.lib import colors
from reportlab.pdfbase.pdfmetrics import stringWidth
from reportlab.pdfgen.canvas import Canvas

from fonts import BOLD, REGULAR
from renderers import RenderError
from slides import (MAX_SLIDES, MAX_DECK_TITLE_CHARS, MAX_SLIDE_TITLE_CHARS, MAX_BULLETS,
                    MAX_BULLET_CHARS, MAX_DECK_CHARS, PAGE_WIDTH, PAGE_HEIGHT, MARGIN,
                    TITLE_SIZE, TITLE_LEADING, BODY_SIZE, BODY_LEADING, BULLET_GAP,
                    TEXT_WIDTH, MEASURE_WIDTH, TextBlock, SlideLayout, _wrap, _validate_text)
from visual_specs import validate_chart, validate_table

THEMES = {
    "droplet": {"background": "FFFFFF", "ink": "173042", "accent": "007E87", "muted": "61717D", "panel": "EDF4F5"},
    "light": {"background": "FFFFFF", "ink": "202A35", "accent": "2962B9", "muted": "566575", "panel": "EEF2F7"},
    "dark": {"background": "12212A", "ink": "F4F7FA", "accent": "51D1CB", "muted": "ADC1CD", "panel": "223B49"},
}
SERIES_COLORS = ["007E87", "4368C4", "BE5D38", "825AA4"]
def _palette(theme):
    return ["51D1CB", "7C9FFA", "F19A78", "C59BE8"] if theme["background"] == "12212A" else SERIES_COLORS
BOTTOM = PAGE_HEIGHT - MARGIN - 18


def _bullet_blocks(raw, index, top, x=MARGIN, width=TEXT_WIDTH, size=BODY_SIZE, leading=BODY_LEADING):
    if not isinstance(raw, list):
        raise RenderError(f"slide {index} bullets must be an array")
    if len(raw) > MAX_BULLETS:
        raise RenderError(f"slide {index} has too many bullets (max {MAX_BULLETS})")
    result = []
    for n, bullet in enumerate(raw, 1):
        text = _validate_text(bullet, MAX_BULLET_CHARS, f"slide {index} bullet {n}", required=True)
        lines = _wrap(text, width * .90 - 26, size)
        if top + len(lines) * leading > BOTTOM:
            raise RenderError(f"slide {index} content does not fit; split it into more slides")
        result.append(TextBlock(lines, x + 26, top, size, leading, bullet=True, width=width - 26))
        top += len(lines) * leading + BULLET_GAP
    return result


def _table(spec, index, top):
    rows, heights = [], []
    width = TEXT_WIDTH / len(spec["headers"])
    for n, row in enumerate([spec["headers"], *spec["rows"]]):
        wrapped = [_wrap(cell, width * .9 - 16, 16, bold=n == 0) for cell in row]
        rows.append(wrapped)
        heights.append(max(map(len, wrapped)) * 22 + 12)
    if top + sum(heights) > BOTTOM:
        raise RenderError(f"slide {index} table does not fit; split or shorten it")
    return {"rows": rows, "heights": heights, "top": top}


def _chart(spec, index, top):
    height = BOTTOM - top
    if height < 230:
        raise RenderError(f"slide {index} chart does not fit; shorten title or subtitle")
    pie = spec["kind"] == "pie"
    legend = spec["labels"] if pie else [entry["name"] for entry in spec["series"]]
    legend_width = 275 if pie else TEXT_WIDTH / len(legend) - 24
    if any(stringWidth(text, REGULAR, 14) > legend_width * .9 for text in legend):
        raise RenderError(f"slide {index} chart legend does not fit; shorten names")
    if pie and len(legend) * 24 > height:
        raise RenderError(f"slide {index} pie legend does not fit; use fewer categories")
    label_width = (TEXT_WIDTH - 70) / len(spec["labels"]) - 10
    labels = [] if pie else [_wrap(text, label_width * .9, 14) for text in spec["labels"]]
    if any(len(lines) > 2 for lines in labels):
        raise RenderError(f"slide {index} chart labels do not fit; shorten labels")
    values = [v for entry in spec["series"] for v in entry["values"]]
    minimum, maximum = min(0, min(values)), max(0, max(values))
    if minimum == maximum:
        maximum = minimum + 1
    return {"spec": spec, "top": top, "height": height, "legend": legend,
            "labels": labels, "minimum": minimum, "maximum": maximum}


def layout_deck(title: str, specs: list[dict[str, Any]]) -> list[SlideLayout]:
    _validate_text(title, MAX_DECK_TITLE_CHARS, "deck title")
    if not isinstance(specs, list) or not specs:
        raise RenderError("at least one slide is required")
    if len(specs) > MAX_SLIDES:
        raise RenderError(f"too many slides (max {MAX_SLIDES})")
    total, result = len(title), []
    for index, spec in enumerate(specs, 1):
        if not isinstance(spec, dict) or set(spec) - {"title", "bullets", "subtitle", "columns", "table", "chart", "notes"}:
            raise RenderError(f"slide {index} accepts only title and bullets, subtitle, columns, table, chart and notes")
        heading = _validate_text(spec.get("title"), MAX_SLIDE_TITLE_CHARS, f"slide {index} title", required=True)
        subtitle = _validate_text(spec.get("subtitle", ""), 300, f"slide {index} subtitle")
        notes = _validate_text(spec.get("notes", ""), 4000, f"slide {index} notes")
        raw, columns = spec.get("bullets", []), spec.get("columns", [])
        if not isinstance(raw, list) or not isinstance(columns, list):
            raise RenderError(f"slide {index} bullets and columns must be an array")
        if sum([bool(raw), bool(columns), spec.get("table") is not None, spec.get("chart") is not None]) > 1:
            raise RenderError(f"slide {index} must choose one layout: bullets, columns, table or chart")
        lines = _wrap(heading, MEASURE_WIDTH, TITLE_SIZE, True)
        if len(lines) > 3:
            raise RenderError(f"slide {index} title does not fit; shorten it")
        heading_block = TextBlock(lines, MARGIN, MARGIN, TITLE_SIZE, TITLE_LEADING, True)
        cursor = MARGIN + len(lines) * TITLE_LEADING + 30
        subtitle_block = None
        if subtitle:
            lines = _wrap(subtitle, MEASURE_WIDTH, 18)
            if len(lines) > 2:
                raise RenderError(f"slide {index} subtitle does not fit; shorten it")
            subtitle_block = TextBlock(lines, MARGIN, cursor, 18, 25)
            cursor += len(lines) * 25 + 18
        blocks, headings, table, chart = [], [], None, None
        text_count = len(heading) + len(subtitle) + len(notes)
        if columns:
            if len(columns) != 2:
                raise RenderError(f"slide {index} requires exactly two columns")
            width = (TEXT_WIDTH - 40) / 2
            for n, column in enumerate(columns):
                if not isinstance(column, dict) or set(column) - {"title", "bullets"}:
                    raise RenderError(f"slide {index} columns accept title and bullets")
                column_title = _validate_text(column.get("title", ""), 100, f"slide {index} column title")
                top, x = cursor, MARGIN + n * (width + 40)
                if column_title:
                    lines = _wrap(column_title, width * .90, 20, True)
                    if len(lines) > 2:
                        raise RenderError(f"slide {index} column title does not fit")
                    headings.append(TextBlock(lines, x, top, 20, 28, True, width=width))
                    top += len(lines) * 28 + 12
                items = column.get("bullets", [])
                blocks.extend(_bullet_blocks(items, index, top, x, width, 19, 26))
                text_count += len(column_title) + sum(map(len, items))
        elif spec.get("table") is not None:
            data = validate_table(spec["table"], f"slide {index} table")
            table = _table(data, index, cursor)
            text_count += sum(len(cell) for row in [data["headers"], *data["rows"]] for cell in row)
        elif spec.get("chart") is not None:
            data = validate_chart(spec["chart"], f"slide {index} chart")
            chart = _chart(data, index, cursor)
            text_count += sum(map(len, data["labels"])) + sum(len(e["name"]) for e in data["series"])
        else:
            blocks = _bullet_blocks(raw, index, cursor)
            text_count += sum(map(len, raw))
        total += text_count
        if total > MAX_DECK_CHARS:
            raise RenderError(f"deck text is too long (max {MAX_DECK_CHARS} characters)")
        result.append(SlideLayout(heading_block, blocks, subtitle_block, headings, table, chart, notes))
    return result


def _color(value):
    return colors.HexColor("#" + value)


def _number(value):
    for scale, suffix in [(1e12, "T"), (1e9, "B"), (1e6, "M"), (1e3, "K")]:
        if abs(value) >= scale:
            return f"{value / scale:g}{suffix}"
    return f"{value:.3g}"


def _pdf_chart(canvas, chart, theme):
    spec, top, height = chart["spec"], chart["top"], chart["height"]
    canvas.setFont(REGULAR, 14)
    palette = _palette(theme)
    if spec["kind"] == "pie":
        diameter = min(height - 12, 350)
        left, bottom, angle = MARGIN + 70, PAGE_HEIGHT - top - diameter, 90.0
        values = spec["series"][0]["values"]
        for n, value in enumerate(values):
            extent = value / sum(values) * 360
            canvas.setFillColor(_color(palette[n % 4]))
            canvas.wedge(left, bottom, left + diameter, bottom + diameter, angle, extent, stroke=0, fill=1)
            angle += extent
            y = PAGE_HEIGHT - top - n * 24 - 16
            canvas.rect(MARGIN + 520, y, 10, 10, stroke=0, fill=1)
            canvas.setFillColor(_color(theme["ink"]))
            canvas.drawString(MARGIN + 540, y, chart["legend"][n])
        return
    left, width = MARGIN + 65, TEXT_WIDTH - 70
    bottom, plot_height = PAGE_HEIGHT - top - height + 50, height - 95
    step = width / len(spec["labels"])
    def ordinate(value):
        return bottom + (value - chart["minimum"]) / (chart["maximum"] - chart["minimum"]) * plot_height
    for n in range(5):
        value = chart["minimum"] + (chart["maximum"] - chart["minimum"]) * n / 4
        y = ordinate(value)
        canvas.setStrokeColor(_color(theme["panel"]))
        canvas.line(left, y, left + width, y)
        canvas.setFillColor(_color(theme["muted"]))
        canvas.drawRightString(left - 8, y - 4, _number(value))
    for n, lines in enumerate(chart["labels"]):
        for row, line in enumerate(lines):
            canvas.setFillColor(_color(theme["ink"]))
            canvas.drawCentredString(left + (n + .5) * step, bottom - 23 - row * 18, line)
    for n, entry in enumerate(spec["series"]):
        color = _color(palette[n])
        legend_x = MARGIN + n * TEXT_WIDTH / len(spec["series"])
        canvas.setFillColor(color)
        canvas.rect(legend_x, PAGE_HEIGHT - top - 15, 10, 10, stroke=0, fill=1)
        canvas.setFillColor(_color(theme["ink"]))
        canvas.drawString(legend_x + 17, PAGE_HEIGHT - top - 15, entry["name"])
        canvas.setFillColor(color)
        if spec["kind"] == "bar":
            bar_width = step * .72 / len(spec["series"])
            for column, value in enumerate(entry["values"]):
                x = left + column * step + step * .14 + n * bar_width
                zero, y = ordinate(0), ordinate(value)
                canvas.rect(x, min(zero, y), bar_width * .90, abs(y - zero), stroke=0, fill=1)
        else:
            canvas.setStrokeColor(color)
            canvas.setLineWidth(2)
            points = [(left + (column + .5) * step, ordinate(value)) for column, value in enumerate(entry["values"])]
            for a, b in zip(points, points[1:]):
                canvas.line(*a, *b)
            for x, y in points:
                canvas.circle(x, y, 3, stroke=0, fill=1)


def _text_blocks(layout):
    return [layout.title, *([layout.subtitle] if layout.subtitle else []), *(layout.headings or []), *layout.bullets]


def render_pdf(title, layouts, theme):
    buf = io.BytesIO()
    canvas = Canvas(buf, pagesize=(PAGE_WIDTH, PAGE_HEIGHT), pageCompression=1)
    canvas.setTitle(title or "Slide deck")
    canvas.setAuthor("Droplet")
    for index, layout in enumerate(layouts, 1):
        canvas.setFillColor(_color(theme["background"]))
        canvas.rect(0, 0, PAGE_WIDTH, PAGE_HEIGHT, stroke=0, fill=1)
        canvas.setFillColor(_color(theme["accent"]))
        canvas.rect(MARGIN, PAGE_HEIGHT - 34, 54, 4, stroke=0, fill=1)
        for block in _text_blocks(layout):
            canvas.setFillColor(_color(theme["muted"] if block is layout.subtitle else theme["ink"]))
            canvas.setFont(BOLD if block.bold else REGULAR, block.font_size)
            for row, line in enumerate(block.lines):
                canvas.drawString(block.x, PAGE_HEIGHT - block.top - block.font_size - row * block.leading, line)
            if block.bullet:
                canvas.setFillColor(_color(theme["accent"]))
                canvas.circle(block.x - 20, PAGE_HEIGHT - block.top - block.font_size * .62, 3.5, stroke=0, fill=1)
        if layout.table:
            table, cursor = layout.table, layout.table["top"]
            width = TEXT_WIDTH / len(table["rows"][0])
            for n, (row, height) in enumerate(zip(table["rows"], table["heights"])):
                canvas.setFillColor(_color(theme["panel"] if n == 0 else theme["background"]))
                canvas.rect(MARGIN, PAGE_HEIGHT - cursor - height, TEXT_WIDTH, height, stroke=0, fill=1)
                canvas.setFillColor(_color(theme["ink"]))
                canvas.setFont(BOLD if n == 0 else REGULAR, 16)
                for column, lines in enumerate(row):
                    for line_index, line in enumerate(lines):
                        canvas.drawString(MARGIN + column * width + 8, PAGE_HEIGHT - cursor - 22 - line_index * 22, line)
                canvas.setStrokeColor(_color(theme["muted"]))
                canvas.setLineWidth(.5)
                canvas.line(MARGIN, PAGE_HEIGHT - cursor - height, MARGIN + TEXT_WIDTH, PAGE_HEIGHT - cursor - height)
                cursor += height
        if layout.chart:
            _pdf_chart(canvas, layout.chart, theme)
        if layout.notes:
            canvas.textAnnotation(layout.notes, Rect=(MARGIN, 20, MARGIN + 18, 38))
        canvas.setFillColor(_color(theme["muted"]))
        canvas.setFont(REGULAR, 10)
        canvas.drawRightString(PAGE_WIDTH - MARGIN, 28, f"{index} / {len(layouts)}")
        canvas.showPage()
    canvas.save()
    return buf.getvalue()


def render_pptx(title, layouts, theme):
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_AUTO_SIZE, MSO_ANCHOR, PP_ALIGN
    from pptx.util import Pt

    def paint(shape, color):
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(color)
        shape.line.fill.background()
        shape.shadow.inherit = False

    def frame_text(frame, lines, size, leading, bold=False, color=None):
        frame.clear()
        frame.margin_top = frame.margin_bottom = frame.margin_left = frame.margin_right = 0
        frame.auto_size, frame.word_wrap, frame.vertical_anchor = MSO_AUTO_SIZE.NONE, False, MSO_ANCHOR.TOP
        for n, line in enumerate(lines):
            p = frame.paragraphs[0] if n == 0 else frame.add_paragraph()
            p.text, p.font.name, p.font.size, p.font.bold = line, "Arial", Pt(size), bold
            p.font.color.rgb = RGBColor.from_string(color or theme["ink"])
            p.line_spacing, p.space_before, p.space_after = Pt(leading), Pt(0), Pt(0)

    deck = Presentation()
    deck.slide_width, deck.slide_height = Pt(PAGE_WIDTH), Pt(PAGE_HEIGHT)
    deck.core_properties.title, deck.core_properties.author = title or "Slide deck", "Droplet"
    for index, layout in enumerate(layouts, 1):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string(theme["background"])
        paint(slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Pt(MARGIN), Pt(30), Pt(54), Pt(4)), theme["accent"])
        for block in _text_blocks(layout):
            box = slide.shapes.add_textbox(Pt(block.x), Pt(block.top), Pt(block.width), Pt(len(block.lines) * block.leading + 8))
            frame_text(box.text_frame, block.lines, block.font_size, block.leading, block.bold,
                       theme["muted"] if block is layout.subtitle else None)
            if block.bullet:
                paint(slide.shapes.add_shape(MSO_SHAPE.OVAL, Pt(block.x - 23.5),
                      Pt(block.top + block.font_size * .62 - 3.5), Pt(7), Pt(7)), theme["accent"])
        if layout.table:
            data = layout.table
            table = slide.shapes.add_table(len(data["rows"]), len(data["rows"][0]), Pt(MARGIN), Pt(data["top"]),
                                           Pt(TEXT_WIDTH), Pt(sum(data["heights"]))).table
            for row_index, (row, height) in enumerate(zip(data["rows"], data["heights"])):
                table.rows[row_index].height = Pt(height)
                for column, lines in enumerate(row):
                    cell = table.cell(row_index, column)
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = RGBColor.from_string(theme["panel"] if row_index == 0 else theme["background"])
                    frame_text(cell.text_frame, lines, 16, 22, row_index == 0)
                    cell.margin_top, cell.margin_bottom, cell.margin_left, cell.margin_right = Pt(6), Pt(6), Pt(8), Pt(8)
        if layout.chart:
            bounds, spec = layout.chart, layout.chart["spec"]
            data = CategoryChartData()
            data.categories = ["\n".join(lines) for lines in bounds["labels"]] if bounds["labels"] else spec["labels"]
            for entry in spec["series"]:
                data.add_series(entry["name"], entry["values"])
            kind = {"bar": XL_CHART_TYPE.COLUMN_CLUSTERED, "line": XL_CHART_TYPE.LINE_MARKERS, "pie": XL_CHART_TYPE.PIE}[spec["kind"]]
            chart = slide.shapes.add_chart(kind, Pt(MARGIN), Pt(bounds["top"]), Pt(TEXT_WIDTH), Pt(bounds["height"]), data).chart
            chart.has_title = False
            for auto_title in chart._chartSpace.xpath("./c:chart/c:autoTitleDeleted"):
                auto_title.set("val", "1")
            # python-pptx 1.0.2's built-in chart templates contain signed axis
            # ids, while DrawingML's ST_UnsignedInt requires uint32. Preserve
            # reference identity while making the generated OOXML valid.
            for axis_id in chart._chartSpace.xpath(".//c:axId | .//c:crossAx"):
                axis_id.set("val", str(int(axis_id.get("val")) & 0xFFFFFFFF))
            chart.has_legend = True
            chart.legend.position = XL_LEGEND_POSITION.RIGHT if spec["kind"] == "pie" else XL_LEGEND_POSITION.TOP
            chart.legend.include_in_layout = False
            chart.legend.font.name, chart.legend.font.size = "Arial", Pt(14)
            chart.legend.font.color.rgb = RGBColor.from_string(theme["ink"])
            if spec["kind"] != "pie":
                for axis in [chart.category_axis, chart.value_axis]:
                    axis.tick_labels.font.name, axis.tick_labels.font.size = "Arial", Pt(14)
                    axis.tick_labels.font.color.rgb = RGBColor.from_string(theme["ink"])
                chart.value_axis.minimum_scale, chart.value_axis.maximum_scale = bounds["minimum"], bounds["maximum"]
            palette = _palette(theme)
            for n, series in enumerate(chart.series):
                series.format.fill.solid()
                series.format.fill.fore_color.rgb = RGBColor.from_string(palette[n])
                series.format.line.color.rgb = RGBColor.from_string(palette[n])
            if spec["kind"] == "pie":
                for n, point in enumerate(chart.series[0].points):
                    point.format.fill.solid()
                    point.format.fill.fore_color.rgb = RGBColor.from_string(palette[n % 4])
        if layout.notes:
            slide.notes_slide.notes_text_frame.text = layout.notes
        footer = slide.shapes.add_textbox(Pt(PAGE_WIDTH - MARGIN - 120), Pt(PAGE_HEIGHT - 40), Pt(120), Pt(20))
        frame_text(footer.text_frame, [f"{index} / {len(layouts)}"], 10, 14, color=theme["muted"])
        footer.text_frame.paragraphs[0].alignment = PP_ALIGN.RIGHT
    buf = io.BytesIO()
    deck.save(buf)
    return buf.getvalue()
