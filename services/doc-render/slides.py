"""Stateless slide writer: shared measured layout for PDF and editable PPTX.

No images, links, HTML, file paths, or external resources enter this spec.
Rather than shrink or discard crowded text, reject it with a slide number so
the caller can split that slide. PDF uses the built-in Helvetica font. The
supported character set is explicit so its missing-glyph box is never shipped.
"""

from __future__ import annotations

import io
from dataclasses import dataclass
from typing import Any

from reportlab.lib import colors
from reportlab.pdfbase.pdfmetrics import stringWidth
from reportlab.pdfgen.canvas import Canvas

from renderers import RenderError

MAX_SLIDES = 60
MAX_DECK_TITLE_CHARS = 255  # OOXML core-properties hard limit; never truncate.
MAX_SLIDE_TITLE_CHARS = 160
MAX_BULLETS = 8
MAX_BULLET_CHARS = 500
MAX_DECK_CHARS = 50_000
PAGE_WIDTH, PAGE_HEIGHT = 960, 540  # points, 16:9
MARGIN = 60
TITLE_SIZE, TITLE_LEADING = 32, 40
BODY_SIZE, BODY_LEADING = 22, 30
BULLET_GAP = 10
TEXT_WIDTH = PAGE_WIDTH - MARGIN * 2
# Arial in PowerPoint is close to Helvetica. Leave width and height headroom
# for font substitution on the viewer's machine; neither format auto-shrinks.
MEASURE_WIDTH = TEXT_WIDTH * 0.90
INK = "173042"
ACCENT = "007E87"
MUTED = "61717D"


@dataclass(frozen=True)
class TextBlock:
    lines: list[str]
    x: float
    top: float
    font_size: int
    leading: int
    bold: bool = False


@dataclass(frozen=True)
class SlideLayout:
    title: TextBlock
    bullets: list[TextBlock]


def _validate_text(text: Any, limit: int, location: str, *, required: bool = False) -> str:
    if not isinstance(text, str):
        raise RenderError(f"{location} must be a string")
    if len(text) > limit:
        raise RenderError(f"{location} is too long (max {limit} characters)")
    if required and not text.strip():
        raise RenderError(f"{location} must not be empty")
    for char in text:
        if char == "\n":
            continue
        if ord(char) < 32 or 127 <= ord(char) < 160:
            raise RenderError(f"{location} contains an unsupported control character")
        try:
            char.encode("cp1252")
        except UnicodeEncodeError:
            raise RenderError(
                f"{location} contains a character unsupported by the slide font "
                f"(U+{ord(char):04X}); use Latin text"
            ) from None
    return text


def _wrap(text: str, width: float, size: int, bold: bool = False) -> list[str]:
    """Wrap plain text, including explicit breaks and long unbroken words."""
    font = "Helvetica-Bold" if bold else "Helvetica"
    lines: list[str] = []
    for paragraph in text.split("\n"):
        current = ""
        for word in paragraph.split():
            candidate = f"{current} {word}" if current else word
            if stringWidth(candidate, font, size) <= width:
                current = candidate
                continue
            if current:
                lines.append(current)
                current = ""
            # URLs and identifiers are still text, never fetch targets. Break
            # long words at a character boundary instead of clipping them.
            for char in word:
                if stringWidth(current + char, font, size) > width and current:
                    lines.append(current)
                    current = ""
                current += char
        lines.append(current)
    return lines


def _layout(title: str, slides: list[dict[str, Any]]) -> list[SlideLayout]:
    _validate_text(title, MAX_DECK_TITLE_CHARS, "deck title")
    if not isinstance(slides, list) or not slides:
        raise RenderError("at least one slide is required")
    if len(slides) > MAX_SLIDES:
        raise RenderError(f"too many slides (max {MAX_SLIDES})")
    total_chars = len(title)
    layouts = []
    for index, spec in enumerate(slides, start=1):
        if not isinstance(spec, dict):
            raise RenderError(f"slide {index} must be an object")
        if set(spec) - {"title", "bullets"}:
            raise RenderError(f"slide {index} accepts only title and bullets")
        heading = _validate_text(
            spec.get("title"), MAX_SLIDE_TITLE_CHARS, f"slide {index} title", required=True
        )
        bullets = spec.get("bullets", [])
        if not isinstance(bullets, list):
            raise RenderError(f"slide {index} bullets must be an array")
        if len(bullets) > MAX_BULLETS:
            raise RenderError(f"slide {index} has too many bullets (max {MAX_BULLETS})")
        bullets = [
            _validate_text(bullet, MAX_BULLET_CHARS, f"slide {index} bullet {n}", required=True)
            for n, bullet in enumerate(bullets, start=1)
        ]
        total_chars += len(heading) + sum(map(len, bullets))
        if total_chars > MAX_DECK_CHARS:
            raise RenderError(f"deck text is too long (max {MAX_DECK_CHARS} characters)")

        heading_lines = _wrap(heading, MEASURE_WIDTH, TITLE_SIZE, bold=True)
        if len(heading_lines) > 3:
            raise RenderError(f"slide {index} title does not fit; shorten it")
        heading_block = TextBlock(heading_lines, MARGIN, MARGIN, TITLE_SIZE, TITLE_LEADING, True)
        cursor = MARGIN + len(heading_lines) * TITLE_LEADING + 30
        bullet_blocks = []
        for bullet in bullets:
            lines = _wrap(bullet, MEASURE_WIDTH - 26, BODY_SIZE)
            height = len(lines) * BODY_LEADING
            if cursor + height > PAGE_HEIGHT - MARGIN - 18:
                raise RenderError(f"slide {index} content does not fit; split it into more slides")
            bullet_blocks.append(TextBlock(lines, MARGIN + 26, cursor, BODY_SIZE, BODY_LEADING))
            cursor += height + BULLET_GAP
        layouts.append(SlideLayout(heading_block, bullet_blocks))
    return layouts


def render_deck(title: str, slides: list[dict[str, Any]], format: str) -> bytes:
    if format not in ("pdf", "pptx"):
        raise RenderError("slide decks require pdf or pptx")
    layouts = _layout(title, slides)
    return _pdf(title, layouts) if format == "pdf" else _pptx(title, layouts)


def _pdf(title: str, layouts: list[SlideLayout]) -> bytes:
    buf = io.BytesIO()
    canvas = Canvas(buf, pagesize=(PAGE_WIDTH, PAGE_HEIGHT), pageCompression=1)
    canvas.setTitle(title or "Slide deck")
    canvas.setAuthor("Droplet")
    for index, layout in enumerate(layouts, start=1):
        canvas.setFillColor(colors.HexColor(f"#{ACCENT}"))
        canvas.rect(MARGIN, PAGE_HEIGHT - 34, 54, 4, stroke=0, fill=1)
        for block in [layout.title, *layout.bullets]:
            canvas.setFillColor(colors.HexColor(f"#{INK}"))
            canvas.setFont("Helvetica-Bold" if block.bold else "Helvetica", block.font_size)
            for line_index, line in enumerate(block.lines):
                canvas.drawString(
                    block.x,
                    PAGE_HEIGHT - block.top - block.font_size - line_index * block.leading,
                    line,
                )
            if not block.bold:
                canvas.setFillColor(colors.HexColor(f"#{ACCENT}"))
                canvas.circle(MARGIN + 6, PAGE_HEIGHT - block.top - BODY_SIZE * 0.62, 3.5, stroke=0, fill=1)
        canvas.setFillColor(colors.HexColor(f"#{MUTED}"))
        canvas.setFont("Helvetica", 10)
        canvas.drawRightString(PAGE_WIDTH - MARGIN, 28, f"{index} / {len(layouts)}")
        canvas.showPage()
    canvas.save()
    return buf.getvalue()


def _pptx(title: str, layouts: list[SlideLayout]) -> bytes:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_AUTO_SIZE, MSO_ANCHOR, PP_ALIGN
    from pptx.util import Pt

    deck = Presentation()
    deck.slide_width, deck.slide_height = Pt(PAGE_WIDTH), Pt(PAGE_HEIGHT)
    deck.core_properties.title = title or "Slide deck"
    deck.core_properties.author = "Droplet"
    for index, layout in enumerate(layouts, start=1):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255)
        accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Pt(MARGIN), Pt(30), Pt(54), Pt(4))
        accent.fill.solid()
        accent.fill.fore_color.rgb = RGBColor.from_string(ACCENT)
        accent.line.fill.background()
        accent.shadow.inherit = False
        for block in [layout.title, *layout.bullets]:
            box = slide.shapes.add_textbox(
                Pt(block.x), Pt(block.top), Pt(TEXT_WIDTH - (block.x - MARGIN)),
                Pt(len(block.lines) * block.leading + 8),
            )
            frame = box.text_frame
            frame.clear()
            frame.margin_top = frame.margin_bottom = frame.margin_left = frame.margin_right = 0
            frame.auto_size = MSO_AUTO_SIZE.NONE
            frame.word_wrap = False  # shared measured breaks, with generous width headroom
            frame.vertical_anchor = MSO_ANCHOR.TOP
            for line_index, line in enumerate(block.lines):
                paragraph = frame.paragraphs[0] if line_index == 0 else frame.add_paragraph()
                paragraph.text = line
                paragraph.font.name = "Arial"
                paragraph.font.size = Pt(block.font_size)
                paragraph.font.bold = block.bold
                paragraph.font.color.rgb = RGBColor.from_string(INK)
                paragraph.line_spacing = Pt(block.leading)
                paragraph.space_before = paragraph.space_after = Pt(0)
            if not block.bold:
                dot = slide.shapes.add_shape(
                    MSO_SHAPE.OVAL, Pt(MARGIN + 2.5), Pt(block.top + BODY_SIZE * 0.62 - 3.5), Pt(7), Pt(7)
                )
                dot.fill.solid()
                dot.fill.fore_color.rgb = RGBColor.from_string(ACCENT)
                dot.line.fill.background()
                dot.shadow.inherit = False
        footer = slide.shapes.add_textbox(Pt(PAGE_WIDTH - MARGIN - 120), Pt(PAGE_HEIGHT - 40), Pt(120), Pt(20))
        footer.text_frame.margin_top = footer.text_frame.margin_bottom = 0
        paragraph = footer.text_frame.paragraphs[0]
        paragraph.text = f"{index} / {len(layouts)}"
        paragraph.alignment = PP_ALIGN.RIGHT
        paragraph.font.name, paragraph.font.size = "Arial", Pt(10)
        paragraph.font.color.rgb = RGBColor.from_string(MUTED)
    buf = io.BytesIO()
    deck.save(buf)
    return buf.getvalue()
