"""Real Office packages: inspect/revise preservation and hostile ZIP/XML guards."""
import base64
import io
import zipfile

import pytest
from lxml import etree as ET
from docx import Document
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill
from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches, Pt

from office_files import OfficeError, inspect_office, revise_office, W, A, S, P, R


def save(document):
    stream = io.BytesIO(); document.save(stream); return stream.getvalue()


def parts(raw):
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


def package(data):
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, body in data.items(): archive.writestr(name, body)
    return stream.getvalue()


def png():
    stream = io.BytesIO(); Image.new("RGB", (20, 20), "blue").save(stream, "PNG"); stream.seek(0); return stream


@pytest.fixture
def docx():
    document = Document()
    paragraph = document.add_paragraph()
    run = paragraph.add_run("Original "); run.bold = True
    paragraph.add_run("mixed formatting").italic = True
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Table first"; table.cell(0, 1).text = "Table second"
    document.sections[0].header.paragraphs[0].text = "Header"
    document.sections[0].footer.paragraphs[0].text = "Footer"
    document.add_picture(png(), width=Inches(0.2))
    return save(document)


@pytest.fixture
def xlsx():
    workbook = Workbook(); sheet = workbook.active; sheet.title = "Revenue"
    sheet.append(["Quarter", "Revenue", "Tax"])
    sheet.append(["Q1", 120, "=B2*0.2"]); sheet.append(["Q2", 180, "=B3*0.2"])
    sheet["B2"].font = Font(bold=True, color="123456")
    sheet["B2"].fill = PatternFill("solid", fgColor="EEEEEE"); sheet["B2"].number_format = "$#,##0.00"
    chart = BarChart(); chart.title = "Revenue"; chart.add_data(Reference(sheet, min_col=2, min_row=1, max_row=3), titles_from_data=True); chart.set_categories(Reference(sheet, min_col=1, min_row=2, max_row=3)); sheet.add_chart(chart, "E2")
    sheet.add_image(__import__("openpyxl").drawing.image.Image(png()), "L1")
    other = workbook.create_sheet("Unchanged"); other["A1"] = "Keep me"
    return save(workbook)


@pytest.fixture
def pptx():
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
    paragraph = box.text_frame.paragraphs[0]
    run = paragraph.add_run(); run.text = "Original "; run.font.bold = True; run.font.size = Pt(24)
    paragraph.add_run().text = "slide title"
    slide.shapes.add_picture(png(), Inches(7), Inches(1), width=Inches(0.2))
    chart = CategoryChartData(); chart.categories = ["A", "B"]; chart.add_series("Sales", [12, 18])
    slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(3), Inches(5), Inches(3), chart)
    slide.notes_slide.notes_text_frame.text = "Speaker notes"
    other = presentation.slides.add_slide(presentation.slide_layouts[6])
    other.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text = "Second slide"
    return save(presentation)


def test_docx_inspects_body_tables_header_footer_and_revises_exact_paragraph(docx):
    inspected = inspect_office(docx, "docx")
    texts = [p["text"] for p in inspected["paragraphs"]]
    assert all(t in texts for t in ["Original mixed formatting", "Table first", "Table second", "Header", "Footer"])
    body = next(p for p in inspected["paragraphs"] if p["text"].startswith("Original"))
    table = next(p for p in inspected["paragraphs"] if p["text"] == "Table second")
    revision = revise_office(docx, "docx", {"text": [{"id": body["id"], "text": "Revised title"}, {"id": table["id"], "text": "Revised cell"}]})
    reopened = Document(io.BytesIO(revision))
    assert reopened.paragraphs[0].text == "Revised title"
    assert reopened.paragraphs[0].runs[0].bold is True
    assert reopened.tables[0].cell(0, 1).text == "Revised cell"
    old, new = parts(docx), parts(revision)
    assert old.keys() == new.keys()
    assert [name for name in old if old[name] != new[name]] == ["word/document.xml"]
    assert Document(io.BytesIO(docx)).paragraphs[0].text == "Original mixed formatting"


def test_xlsx_revises_existing_values_and_preserves_chart_style_image_and_other_sheet(xlsx):
    inspected = inspect_office(xlsx, "xlsx")
    assert [s["name"] for s in inspected["sheets"]] == ["Revenue", "Unchanged"]
    assert next(c for c in inspected["sheets"][0]["cells"] if c["cell"] == "C2")["formula"] == "B2*0.2"
    revision = revise_office(xlsx, "xlsx", {"cells": [{"sheet": "Revenue", "cell": "B2", "value": 250}, {"sheet": "Revenue", "cell": "A2", "value": "=not-a-formula"}]})
    reopened = load_workbook(io.BytesIO(revision))
    assert reopened["Revenue"]["B2"].value == 250
    assert reopened["Revenue"]["B2"].font.bold is True
    assert reopened["Revenue"]["B2"].number_format == "$#,##0.00"
    assert reopened["Revenue"]["A2"].value == "=not-a-formula" and reopened["Revenue"]["A2"].data_type == "s"
    assert reopened["Revenue"]["C2"].value == "=B2*0.2"
    assert reopened.calculation.fullCalcOnLoad and reopened.calculation.forceFullCalc
    assert len(reopened["Revenue"]._charts) == 1 and len(reopened["Revenue"]._images) == 1
    old, new = parts(xlsx), parts(revision)
    assert old["xl/styles.xml"] == new["xl/styles.xml"]
    assert old["xl/worksheets/sheet2.xml"] == new["xl/worksheets/sheet2.xml"]
    assert all(old[n] == new[n] for n in old if n.startswith("xl/media/"))


def test_pptx_preserves_native_chart_embedded_workbook_image_notes_and_other_slide(pptx):
    inspected = inspect_office(pptx, "pptx")
    body = next(p for p in inspected["paragraphs"] if p["text"] == "Original slide title")
    assert any(p["text"] == "Speaker notes" for p in inspected["paragraphs"])
    revision = revise_office(pptx, "pptx", {"text": [{"id": body["id"], "text": "Revised slide title"}]})
    reopened = Presentation(io.BytesIO(revision))
    paragraph = reopened.slides[0].shapes[0].text_frame.paragraphs[0]
    assert paragraph.text == "Revised slide title" and paragraph.runs[0].font.bold is True
    assert reopened.slides[0].notes_slide.notes_text_frame.text == "Speaker notes"
    assert reopened.slides[0].shapes[2].chart.series[0].values == (12.0, 18.0)
    old, new = parts(pptx), parts(revision)
    assert [name for name in old if old[name] != new[name]] == ["ppt/slides/slide1.xml"]
    assert any("embeddings/" in name for name in old)


def test_presentation_inspection_uses_slide_order_not_zip_part_number(pptx):
    data = parts(pptx); root = ET.fromstring(data["ppt/presentation.xml"])
    slides = root.find(f"{{{P}}}sldIdLst"); first = slides[0]; slides.remove(first); slides.append(first)
    data["ppt/presentation.xml"] = ET.tostring(root)
    inspected = inspect_office(package(data), "pptx")
    assert inspected["paragraphs"][0]["text"] == "Second slide"


@pytest.mark.parametrize("format", ["docx", "xlsx", "pptx"])
def test_http_office_requires_auth_and_returns_inspection_or_new_ooxml(client, auth, request, format):
    raw = request.getfixturevalue(format)
    payload = {"format": format, "content_base64": base64.b64encode(raw).decode()}
    assert client.post("/office", json=payload).status_code == 401
    response = client.post("/office", json=payload, headers=auth)
    assert response.status_code == 200 and response.json()["format"] == format
    if format == "xlsx": changes = {"cells": [{"sheet": "Revenue", "cell": "B2", "value": 25}]}
    else:
        target = next(p for p in response.json()["paragraphs"] if p["text"].startswith("Original"))
        changes = {"text": [{"id": target["id"], "text": "Updated"}]}
    revised = client.post("/office", json={**payload, "action": "revise", "changes": changes}, headers=auth)
    assert revised.status_code == 200 and zipfile.is_zipfile(io.BytesIO(revised.content))
    assert revised.content != raw


@pytest.mark.parametrize("attack", ["macro", "ole", "external", "dtd", "symlink", "traversal", "duplicate", "zipbomb", "active-type", "missing-target"])
def test_rejects_hostile_unused_package_parts(docx, attack):
    data = parts(docx)
    if attack == "macro": data["word/vbaProject.bin"] = b"active"
    elif attack == "ole": data["word/embeddings/oleObject1.bin"] = b"object"
    elif attack == "external": data["word/_rels/unused.xml.rels"] = f'<Relationships xmlns="{R}"><Relationship Id="r1" Type="x" Target="https://evil.invalid/a" TargetMode="External"/></Relationships>'.encode()
    elif attack == "dtd": data["word/unused.xml"] = b'<!DOCTYPE x [<!ENTITY secret SYSTEM "file:///etc/passwd">]><x>&secret;</x>'
    elif attack == "traversal": data["../hidden.xml"] = b"<x/>"
    elif attack == "zipbomb": data["word/bomb.bin"] = b"x" * 100_000
    elif attack == "active-type": data["[Content_Types].xml"] = data["[Content_Types].xml"].replace(b"wordprocessingml.document.main", b"wordprocessingml.document.macroEnabled.main")
    elif attack == "missing-target": data["word/_rels/unused.xml.rels"] = f'<Relationships xmlns="{R}"><Relationship Id="r1" Type="x" Target="absent.xml"/></Relationships>'.encode()
    raw = package(data)
    if attack in ("symlink", "duplicate"):
        stream = io.BytesIO(raw)
        with zipfile.ZipFile(stream, "a") as archive:
            if attack == "duplicate":
                with pytest.warns(UserWarning): archive.writestr("word/document.xml", data["word/document.xml"])
            else:
                info = zipfile.ZipInfo("linked"); info.create_system = 3; info.external_attr = 0o120777 << 16; archive.writestr(info, "../secret")
        raw = stream.getvalue()
    with pytest.raises(OfficeError): inspect_office(raw, "docx")
    with pytest.raises(OfficeError): revise_office(raw, "docx", {"text": [{"id": "word/document.xml:p:0", "text": "Changed"}]})


def test_rejects_dangerous_chart_workbook_and_unreferenced_embedded_package(pptx, xlsx):
    data = parts(pptx); name = next(n for n in data if "/embeddings/" in n)
    nested = parts(data[name]); nested["xl/vbaProject.bin"] = b"macro"; data[name] = package(nested)
    with pytest.raises(OfficeError, match="Macros"): inspect_office(package(data), "pptx")
    data = parts(pptx); data["ppt/embeddings/unreferenced.xlsx"] = xlsx
    with pytest.raises(OfficeError, match="Embedded"): inspect_office(package(data), "pptx")


@pytest.mark.parametrize("payload", ['<script>alert(1)</script>', '<image href="https://evil.invalid/image.png"/>', '<rect onload="sendSecret()"/>', '<style>@import "https://evil.invalid/style.css";</style>', '<rect style="fill:url(file:///private)"/>'])
def test_refuses_active_or_external_svg_assets_even_when_unused(docx, payload):
    data = parts(docx); data["word/media/active.svg"] = f'<svg xmlns="http://www.w3.org/2000/svg">{payload}</svg>'.encode()
    with pytest.raises(OfficeError, match="SVG"): inspect_office(package(data), "docx")


@pytest.mark.parametrize("formula", ['WEBSERVICE("https://evil.invalid")', '_xlfn.IMAGE("https://evil.invalid")', 'cmd|\'/c calc\'!A0', "'[external.xlsx]Q1-2026'!A1", 'RTD("provider",,"x")'])
def test_refuses_external_or_dde_formulas_before_recalc_can_be_requested(xlsx, formula):
    data = parts(xlsx); root = ET.fromstring(data["xl/worksheets/sheet1.xml"])
    root.find(f"{{{S}}}sheetData/{{{S}}}row/{{{S}}}c[@r='C2']/{{{S}}}f").text = formula
    data["xl/worksheets/sheet1.xml"] = ET.tostring(root)
    with pytest.raises(OfficeError, match="formulas"): inspect_office(package(data), "xlsx")


def test_local_structured_references_and_formula_string_literals_are_inert(xlsx):
    data = parts(xlsx); root = ET.fromstring(data["xl/worksheets/sheet1.xml"])
    formula = root.find(f"{{{S}}}sheetData/{{{S}}}row/{{{S}}}c[@r='C2']/{{{S}}}f")
    formula.text = 'Table1[Revenue]-Unchanged!A1+IF(A1="WEBSERVICE(|)",1,0)'
    data["xl/worksheets/sheet1.xml"] = ET.tostring(root)
    assert inspect_office(package(data), "xlsx")["format"] == "xlsx"


def test_refuses_dde_word_fields_without_external_relationships(docx):
    data = parts(docx); root = ET.fromstring(data["word/document.xml"])
    paragraph = next(root.iter(f"{{{W}}}p")); run = ET.SubElement(paragraph, f"{{{W}}}r")
    ET.SubElement(run, f"{{{W}}}instrText").text = "DDEAUTO cmd /c calc"
    data["word/document.xml"] = ET.tostring(root)
    with pytest.raises(OfficeError, match="Word fields"): inspect_office(package(data), "docx")


@pytest.mark.parametrize("value", [None, True, False, 1.25, " literal ", "+cmd", "@anything"])
def test_cell_scalar_roundtrip(xlsx, value):
    result = revise_office(xlsx, "xlsx", {"cells": [{"sheet": "Revenue", "cell": "B2", "value": value}]})
    assert load_workbook(io.BytesIO(result))["Revenue"]["B2"].value == value


@pytest.mark.parametrize("changes", [{"cells": []}, {"cells": [{"sheet": "Absent", "cell": "A1", "value": 1}]}, {"cells": [{"sheet": "Revenue", "cell": "Z99", "value": 1}]}, {"cells": [{"sheet": "Revenue", "cell": "XFE1", "value": 1}]}, {"cells": [{"sheet": "Revenue", "cell": "B2", "value": float("inf")}]}, {"cells": [{"sheet": "Revenue", "cell": "B2", "value": "x\0"}]}, {"cells": [{"sheet": "Revenue", "cell": "B2", "value": "\ud800"}]}, {"text": [{"id": "x", "text": "x"}]}])
def test_invalid_cell_operations_fail_without_partial_result(xlsx, changes):
    with pytest.raises(OfficeError): revise_office(xlsx, "xlsx", changes)


def test_duplicate_edits_missing_paragraph_and_paragraph_with_image_are_refused(docx):
    inspection = inspect_office(docx, "docx"); first = inspection["paragraphs"][0]["id"]
    for edits in [[{"id": first, "text": "One"}, {"id": first, "text": "Two"}], [{"id": "missing", "text": "One"}], [{"id": "word/document.xml:p:3", "text": "Image destruction"}]]:
        with pytest.raises(OfficeError): revise_office(docx, "docx", {"text": edits})


def test_shared_array_and_merged_nonanchor_cells_are_refused(xlsx):
    data = parts(xlsx); root = ET.fromstring(data["xl/worksheets/sheet1.xml"])
    formula = root.find(f"{{{S}}}sheetData/{{{S}}}row/{{{S}}}c[@r='C2']/{{{S}}}f")
    formula.set("t", "array"); formula.set("ref", "B2:C3")
    data["xl/worksheets/sheet1.xml"] = ET.tostring(root)
    with pytest.raises(OfficeError, match="formula ranges"): revise_office(package(data), "xlsx", {"cells": [{"sheet": "Revenue", "cell": "B3", "value": 2}]})
    root = ET.fromstring(parts(xlsx)["xl/worksheets/sheet1.xml"])
    ET.SubElement(ET.SubElement(root, f"{{{S}}}mergeCells"), f"{{{S}}}mergeCell", ref="A2:B2")
    data["xl/worksheets/sheet1.xml"] = ET.tostring(root)
    with pytest.raises(OfficeError, match="top-left"): revise_office(package(data), "xlsx", {"cells": [{"sheet": "Revenue", "cell": "B2", "value": 2}]})


def test_inspection_explicitly_marks_bounded_content_and_formula_clipping():
    document = Document()
    for index in range(220): document.add_paragraph(f"Paragraph {index}")
    inspection = inspect_office(save(document), "docx")
    assert inspection["totalItems"] == 220 and inspection["returnedItems"] == 200 and inspection["truncated"] is True
    workbook = Workbook(); workbook.active["A1"] = "=" + "1+" * 2200 + "1"
    inspection = inspect_office(save(workbook), "xlsx")
    cell = inspection["sheets"][0]["cells"][0]
    assert len(cell["formula"]) == 4000 and cell["formulaTruncated"] is True
