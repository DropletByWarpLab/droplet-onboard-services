"""Rich decks preserve native objects, source data and Unicode in both formats."""
from __future__ import annotations

import io
import zipfile

import pytest
from pptx import Presentation
from pypdf import PdfReader

import renderers
from fonts import prepare_text
from visual_specs import validate_chart

TABLE = {"headers": ["Region", "Revenue"], "rows": [["West", "120"], ["East", "180"]]}
CHART = {"kind": "bar", "labels": ["West", "East"], "series": [{"name": "Revenue", "values": [120, 180]}]}


def test_unicode_is_embedded_and_extracted_instead_of_becoming_missing_glyphs():
    spec = [{"title": "Привет κόσμε", "bullets": ["Łódź café — €5", "cafe\u0301"]}]
    reader = PdfReader(io.BytesIO(renderers.render_slide_deck("Unicode", spec, "pdf")))
    text = reader.pages[0].extract_text()
    assert "Привет κόσμε" in text
    assert "Łódź café — €5" in text
    assert "café" in text
    assert any("/FontFile2" in font.get_object()["/FontDescriptor"]
               for font in reader.pages[0]["/Resources"]["/Font"].values()
               if "/FontDescriptor" in font.get_object())


def test_pptx_tables_and_charts_are_native_and_retain_source_values():
    spec = [{"title": "Table", "table": TABLE}, {"title": "Chart", "chart": CHART}]
    payload = renderers.render_slide_deck("Data", spec, "pptx")
    deck = Presentation(io.BytesIO(payload))
    table = next(shape.table for shape in deck.slides[0].shapes if shape.has_table)
    assert [[cell.text for cell in row.cells] for row in table.rows] == [TABLE["headers"], *TABLE["rows"]]
    chart = next(shape.chart for shape in deck.slides[1].shapes if shape.has_chart)
    assert list(chart.series[0].values) == [120, 180]
    assert [c.label for c in chart.plots[0].categories] == ["West", "East"]
    assert all(0 <= int(node.get("val")) <= 0xFFFFFFFF
               for node in chart._chartSpace.xpath(".//c:axId | .//c:crossAx"))
    with zipfile.ZipFile(io.BytesIO(payload)) as package:
        assert any(name.startswith("ppt/embeddings/") and name.endswith(".xlsx") for name in package.namelist())


@pytest.mark.parametrize("kind", ["bar", "line", "pie"])
def test_vector_pdf_chart_preserves_labels_and_has_no_raster_image(kind):
    chart = dict(CHART, kind=kind)
    page = PdfReader(io.BytesIO(renderers.render_slide_deck("Data", [{"title": "Chart", "chart": chart}], "pdf"))).pages[0]
    text = page.extract_text()
    assert "West" in text and "East" in text
    assert not any(obj.get_object().get("/Subtype") == "/Image"
                   for obj in page["/Resources"].get("/XObject", {}).values())


def test_notes_are_preserved_without_adding_presented_pages():
    notes = "Presenter: explain the assumptions. Пример."
    spec = [{"title": "One slide", "notes": notes}]
    deck = Presentation(io.BytesIO(renderers.render_slide_deck("Notes", spec, "pptx")))
    assert len(deck.slides) == 1
    assert deck.slides[0].notes_slide.notes_text_frame.text == notes
    reader = PdfReader(io.BytesIO(renderers.render_slide_deck("Notes", spec, "pdf")))
    assert len(reader.pages) == 1
    annotations = [a.get_object() for a in reader.pages[0]["/Annots"]]
    assert any(a["/Subtype"] == "/Text" and a["/Contents"] == notes for a in annotations)


@pytest.mark.parametrize("theme", ["droplet", "light", "dark"])
@pytest.mark.parametrize("format", ["pdf", "pptx"])
def test_two_column_layout_and_subtitle_keep_both_sections_and_geometry(theme, format):
    spec = [{"title": "Compare", "subtitle": "Two complete sections", "columns": [
        {"title": "Today", "bullets": ["Read files", "Create reports"]},
        {"title": "Next", "bullets": ["Render native charts", "Keep tables editable"]},
    ]}]
    payload = renderers.render_slide_deck("Comparison", spec, format, theme)
    if format == "pdf":
        text = PdfReader(io.BytesIO(payload)).pages[0].extract_text()
    else:
        deck = Presentation(io.BytesIO(payload))
        text = "\n".join(shape.text for shape in deck.slides[0].shapes if shape.has_text_frame)
        for shape in deck.slides[0].shapes:
            assert shape.left + shape.width <= deck.slide_width
            assert shape.top + shape.height <= deck.slide_height
    for expected in ["Two complete sections", "Today", "Next", "Read files", "Keep tables editable"]:
        assert expected in text


@pytest.mark.parametrize("spec,reason", [
    ({"title": "T", "bullets": ["B"], "table": TABLE}, "choose one layout"),
    ({"title": "T", "columns": [{"title": "Only"}]}, "exactly two columns"),
    ({"title": "T", "table": {"headers": ["A", "B"], "rows": [["missing"]]}}, "exactly as many"),
    ({"title": "T", "table": {"headers": ["A"] * 7, "rows": [["a"] * 7]}}, "1-6 headers"),
    ({"title": "T", "table": {"headers": ["A"], "rows": [["a"]] * 11}}, "1-10 rows"),
    ({"title": "T", "table": {"headers": ["A"], "rows": [[3]]}}, "must be strings"),
    ({"title": "T", "table": {"headers": ["A"] * 6, "rows": [["W" * 200] * 6]}}, "table does not fit"),
    ({"title": "T", "columns": [{"bullets": ["x" * 500] * 8}, {"bullets": ["B"]}]}, "content does not fit"),
    ({"title": "T", "subtitle": "W" * 300}, "subtitle does not fit"),
    ({"title": "T", "notes": "x" * 4001}, "notes is too long"),
    ({"title": "T", "chart": dict(CHART, labels=["W" * 60] * 10, series=[{"name": "S", "values": [1] * 10}])}, "labels do not fit"),
])
def test_overcrowded_or_ambiguous_content_is_rejected_with_a_slide_number(spec, reason):
    with pytest.raises(renderers.RenderError, match=reason) as error:
        renderers.render_slide_deck("", [spec], "pdf")
    assert "slide 1" in str(error.value)


@pytest.mark.parametrize("value", [True, "2", float("nan"), float("inf"), 1e13])
def test_chart_values_are_explicit_finite_numbers(value):
    chart = dict(CHART, series=[{"name": "S", "values": [value, 1]}])
    with pytest.raises(renderers.RenderError, match="finite numbers"):
        validate_chart(chart, "chart")


@pytest.mark.parametrize("chart,reason", [
    (dict(CHART, kind="pie", series=[{"name": "S", "values": [-1, 2]}]), "nonnegative"),
    (dict(CHART, kind="pie", series=[{"name": "S", "values": [0, 0]}]), "positive total"),
    (dict(CHART, series=[{"name": "S", "values": [1]}]), "match the category"),
    (dict(CHART, series=[{"name": "S", "values": [1, 2]}] * 5), "1-4 series"),
])
def test_invalid_chart_structure_is_rejected(chart, reason):
    with pytest.raises(renderers.RenderError, match=reason):
        validate_chart(chart, "chart")


@pytest.mark.parametrize("text,reason", [("مرحبا", "shaping"), ("你好", "unsupported by the bundled"), ("bad\x00text", "control")])
def test_unsupported_scripts_and_controls_fail_explicitly(text, reason):
    with pytest.raises(renderers.RenderError, match=reason):
        prepare_text(text, "text")


def test_rich_route_preserves_all_optional_fields_and_rejects_boolean_values(client, auth):
    request = {"format": "pptx", "theme": "dark", "slides": [{"title": "T", "subtitle": "S", "chart": CHART, "notes": "N"}]}
    response = client.post("/render", json=request, headers=auth)
    assert response.status_code == 200
    deck = Presentation(io.BytesIO(response.content))
    assert any(shape.has_chart for shape in deck.slides[0].shapes)
    assert deck.slides[0].notes_slide.notes_text_frame.text == "N"
    request["slides"][0]["chart"] = dict(CHART, series=[{"name": "S", "values": [True, 1]}])
    assert client.post("/render", json=request, headers=auth).status_code == 422
