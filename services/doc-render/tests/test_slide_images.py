from __future__ import annotations

import base64
import io
import zipfile

import pytest
from PIL import Image, PngImagePlugin
from pptx import Presentation
from pypdf import PdfReader

import renderers
import slide_images
from rich_slides import layout_deck


def image_bytes(format="PNG", size=(400, 200), **options):
    buffer = io.BytesIO()
    Image.new("RGB", size, "#007e87").save(buffer, format, **options)
    return buffer.getvalue()


def image_spec(content=None, **values):
    return {"content_base64": base64.b64encode(content or image_bytes()).decode("ascii"), **values}


@pytest.mark.parametrize("format", ["PNG", "JPEG"])
def test_pdf_picture_and_native_pptx_picture_preserve_aspect_and_caption(format):
    spec = [{"title": "Device", "subtitle": "Authorized source", "image": image_spec(image_bytes(format), caption="Droplet prototype", alt="Teal prototype image")}]
    payload = renderers.render_slide_deck("Pictures", spec, "pptx", "dark")
    deck = Presentation(io.BytesIO(payload))
    pictures = [shape for shape in deck.slides[0].shapes if shape.shape_type == 13]
    assert len(pictures) == 1
    picture = pictures[0]
    assert picture.width / picture.height == pytest.approx(2)
    assert picture.left + picture.width <= deck.slide_width
    assert picture.top + picture.height <= deck.slide_height
    assert picture._element.xpath("./p:nvPicPr/p:cNvPr")[0].get("descr") == "Teal prototype image"
    assert "Droplet prototype" in "\n".join(shape.text for shape in deck.slides[0].shapes if shape.has_text_frame)
    with zipfile.ZipFile(io.BytesIO(payload)) as package:
        assert len([name for name in package.namelist() if name.startswith("ppt/media/")]) == 1
        assert all(b'TargetMode="External"' not in package.read(name) for name in package.namelist() if name.endswith(".rels"))
    page = PdfReader(io.BytesIO(renderers.render_slide_deck("Pictures", spec, "pdf"))).pages[0]
    assert "Droplet prototype" in page.extract_text()
    raster = [obj.get_object() for obj in page["/Resources"]["/XObject"].values() if obj.get_object().get("/Subtype") == "/Image"]
    assert len(raster) == 1 and raster[0]["/Width"] == 400 and raster[0]["/Height"] == 200


def test_metadata_and_appended_active_payload_are_removed():
    metadata = PngImagePlugin.PngInfo()
    metadata.add_text("Comment", "private location and <script>payload</script>")
    asset = slide_images.decode_image(image_spec(image_bytes(pnginfo=metadata) + b"<script>appended</script>"), "slide 1 image")
    assert b"private location" not in asset.content and b"script" not in asset.content
    with Image.open(io.BytesIO(asset.content)) as image:
        assert not image.info


def test_compressed_png_metadata_bomb_is_rejected_before_large_expansion():
    metadata = PngImagePlugin.PngInfo()
    for index in range(64):
        metadata.add_text(f"private{index}", "x" * 1_000_000, zip=True)
    content = image_bytes(size=(20, 20), pnginfo=metadata)
    assert len(content) < slide_images.MAX_IMAGE_BYTES
    with pytest.raises(renderers.RenderError, match="invalid or unsafe"):
        slide_images.decode_image(image_spec(content), "slide image")
    assert PngImagePlugin.MAX_TEXT_CHUNK <= 64 * 1024
    assert PngImagePlugin.MAX_TEXT_MEMORY <= 256 * 1024


def test_png_aggregate_text_cap_rejects_many_small_chunks():
    metadata = PngImagePlugin.PngInfo()
    for index in range(6):
        metadata.add_text(f"private{index}", "x" * 60_000, zip=True)
    with pytest.raises(renderers.RenderError, match="invalid or unsafe"):
        slide_images.decode_image(image_spec(image_bytes(size=(20, 20), pnginfo=metadata)), "slide image")


def test_jpeg_exif_orientation_is_applied_and_removed():
    exif = Image.Exif()
    exif[274] = 6
    exif[270] = "secret device metadata"
    asset = slide_images.decode_image(image_spec(image_bytes("JPEG", exif=exif)), "slide image")
    assert (asset.width, asset.height) == (200, 400)
    with Image.open(io.BytesIO(asset.content)) as image:
        assert not image.getexif()


@pytest.mark.parametrize("content", [b"<svg><script>bad</script></svg>", b"%PDF-fake", image_bytes("GIF"), b"\x89PNG\r\n\x1a\ncorrupt", image_bytes("JPEG")[:100]])
def test_active_other_formats_and_truncated_rasters_are_rejected(content):
    with pytest.raises(renderers.RenderError, match="PNG or JPEG|invalid or unsafe"):
        slide_images.decode_image(image_spec(content), "slide image")


@pytest.mark.parametrize("source", [{"url": "https://example.com/p.png"}, {"path": "/image.png"}, {"content_base64": "a!=="}, {"content_base64": "YQ==\n"}])
def test_worker_cannot_fetch_sources_or_accept_malformed_encodings(source):
    with pytest.raises(renderers.RenderError):
        slide_images.decode_image(source, "slide image")


def test_animated_png_is_rejected():
    output = io.BytesIO()
    Image.new("RGB", (20, 20), "red").save(output, "PNG", save_all=True, append_images=[Image.new("RGB", (20, 20), "blue")])
    with pytest.raises(renderers.RenderError, match="single-frame"):
        slide_images.decode_image(image_spec(output.getvalue()), "slide image")


def test_pixel_and_dimension_caps_are_checked_before_decoding(monkeypatch):
    monkeypatch.setattr(slide_images, "MAX_IMAGE_PIXELS", 10)
    with pytest.raises(renderers.RenderError, match="pixel limit"):
        slide_images.decode_image(image_spec(image_bytes(size=(20, 20))), "slide image")
    monkeypatch.setattr(slide_images, "MAX_IMAGE_PIXELS", 16_000_000)
    with pytest.raises(renderers.RenderError, match="dimension"):
        slide_images.decode_image(image_spec(image_bytes(size=(8193, 1))), "slide image")


def test_input_and_reencoded_output_byte_caps(monkeypatch):
    content = image_bytes()
    monkeypatch.setattr(slide_images, "MAX_IMAGE_BYTES", len(content) - 1)
    with pytest.raises(renderers.RenderError, match="3 MiB|canonical"):
        slide_images.decode_image(image_spec(content), "slide image")
    monkeypatch.setattr(slide_images, "MAX_IMAGE_BYTES", len(content) + 20)
    # A small palette PNG expands when converted into canonical RGB pixels.
    output = io.BytesIO()
    Image.new("P", (400, 200)).save(output, "PNG")
    monkeypatch.setattr(slide_images, "MAX_IMAGE_BYTES", len(output.getvalue()) + 20)
    with pytest.raises(renderers.RenderError, match="decoded slide image exceeds"):
        slide_images.decode_image(image_spec(output.getvalue()), "slide image")


def test_image_layout_is_exclusive_caption_is_measured_and_counts_are_bounded(monkeypatch):
    source = image_spec()
    with pytest.raises(renderers.RenderError, match="choose one layout"):
        layout_deck("", [{"title": "T", "image": source, "bullets": ["Conflicting"]}])
    with pytest.raises(renderers.RenderError, match="caption does not fit"):
        layout_deck("", [{"title": "T", "image": image_spec(caption="W" * 300)}])
    with pytest.raises(renderers.RenderError, match="at most 12"):
        layout_deck("", [{"title": "T", "image": source}] * 13)
    import rich_slides
    monkeypatch.setattr(rich_slides, "MAX_DECK_IMAGE_BYTES", len(image_bytes()) + 1)
    with pytest.raises(renderers.RenderError, match="12 MiB"):
        layout_deck("", [{"title": "T", "image": source}] * 2)


def test_total_deck_pixels_are_checked_before_decoding_another_image(monkeypatch):
    import rich_slides
    monkeypatch.setattr(rich_slides, "MAX_DECK_IMAGE_PIXELS", 100_000)
    with pytest.raises(renderers.RenderError, match="total pixel budget"):
        layout_deck("", [{"title": "T", "image": image_spec()}] * 2)
    # A rejected next image never enters the orientation/conversion decode.
    monkeypatch.setattr(slide_images.ImageOps, "exif_transpose", lambda *args, **kwargs: pytest.fail("over-budget image was decoded"))
    with pytest.raises(renderers.RenderError, match="total pixel budget"):
        slide_images.decode_image(image_spec(), "slide image", pixel_budget=1)


def test_authenticated_worker_route_accepts_only_the_hydrated_image_dto(client, auth):
    for format in ["pdf", "pptx"]:
        response = client.post("/render", headers=auth, json={"format": format, "slides": [{"title": "T", "image": image_spec(caption="Kept")}]})
        assert response.status_code == 200
    for source in [{"path": "/p.png"}, {"url": "https://example.com/p.png"}, {"content_base64": base64.b64encode(b"not an image").decode()}]:
        response = client.post("/render", headers=auth, json={"format": "pdf", "slides": [{"title": "T", "image": source}]})
        assert response.status_code in (400, 422)
