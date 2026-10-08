"""Slide decks reopen with intact content and bounded, readable geometry."""

from __future__ import annotations

import io

import pytest
from pptx import Presentation
from pypdf import PdfReader
from reportlab.pdfbase.pdfmetrics import stringWidth

import renderers
import slides

SAMPLE = [
    {"title": "Droplet document creation", "bullets": ["PDF slide decks from chat", "Editable PowerPoint slides"]},
    {"title": "Keep content intact", "bullets": ["Literal <b>text</b> & symbols", "Caf\u00e9 prices: \u20ac5"]},
]


def test_pdf_has_one_widescreen_page_per_slide_and_preserves_text():
    reader = PdfReader(io.BytesIO(renderers.render_slide_deck("Droplet", SAMPLE, "pdf")))
    assert len(reader.pages) == len(SAMPLE)
    assert reader.metadata.title == "Droplet"
    for page, source in zip(reader.pages, SAMPLE):
        assert float(page.mediabox.width) == 960
        assert float(page.mediabox.height) == 540
        text = page.extract_text()
        assert source["title"] in text
        for bullet in source["bullets"]:
            assert bullet in text


def test_pptx_has_editable_native_text_and_shapes_in_widescreen_pages():
    deck = Presentation(io.BytesIO(renderers.render_slide_deck("Droplet", SAMPLE, "pptx")))
    assert deck.slide_width / deck.slide_height == pytest.approx(16 / 9)
    assert len(deck.slides) == len(SAMPLE)
    assert deck.core_properties.title == "Droplet"
    for slide, source in zip(deck.slides, SAMPLE):
        text = "\n".join(shape.text for shape in slide.shapes if shape.has_text_frame)
        assert source["title"] in text
        for bullet in source["bullets"]:
            assert bullet in text
        # Everything can be selected/edited; no screenshot of the whole slide.
        assert not any(shape.shape_type == 13 for shape in slide.shapes)
        for shape in slide.shapes:
            assert shape.left >= 0 and shape.top >= 0
            assert shape.left + shape.width <= deck.slide_width
            assert shape.top + shape.height <= deck.slide_height


@pytest.mark.parametrize("format", ["pdf", "pptx"])
def test_eight_short_bullets_fit_without_shrinking(format):
    spec = [{"title": "Eight items", "bullets": [f"Item {n}" for n in range(8)]}]
    assert renderers.render_slide_deck("", spec, format)


def test_wrapping_long_words_and_explicit_breaks_stays_inside_the_safe_area():
    spec = [{"title": "A longer title " * 6, "bullets": ["A" * 130, "First line\nSecond line"]}]
    layout = slides._layout("", spec)[0]
    assert len(layout.title.lines) > 1
    assert len(layout.bullets[0].lines) > 1
    assert layout.bullets[1].lines == ["First line", "Second line"]
    for block in [layout.title, *layout.bullets]:
        font = "Helvetica-Bold" if block.bold else "Helvetica"
        for line in block.lines:
            assert block.x + stringWidth(line, font, block.font_size) < slides.PAGE_WIDTH - slides.MARGIN
        assert block.top + len(block.lines) * block.leading < slides.PAGE_HEIGHT - slides.MARGIN
    pdf_text = PdfReader(io.BytesIO(renderers.render_slide_deck("", spec, "pdf"))).pages[0].extract_text()
    assert "A" * 130 in pdf_text.replace("\n", "")


@pytest.mark.parametrize("format", ["pdf", "pptx"])
def test_overflow_is_a_numbered_error_instead_of_clipped_content(format):
    spec = [{"title": "Crowded", "bullets": ["Long content " * 30] * 8}]
    with pytest.raises(renderers.RenderError, match="slide 1 content does not fit"):
        renderers.render_slide_deck("", spec, format)


@pytest.mark.parametrize(
    ("spec", "match"),
    [
        ([], "at least one slide"),
        ([{"title": "T"}] * (slides.MAX_SLIDES + 1), "too many slides"),
        ([{"title": "x" * (slides.MAX_SLIDE_TITLE_CHARS + 1)}], "title is too long"),
        ([{"title": "T", "bullets": ["a"] * (slides.MAX_BULLETS + 1)}], "too many bullets"),
        ([{"title": "T", "bullets": ["x" * (slides.MAX_BULLET_CHARS + 1)]}], "bullet 1 is too long"),
        ([{"title": "T", "bullets": [" "]}], "must not be empty"),
        ([{"title": ""}], "must not be empty"),
        ([{"title": "T", "bullets": "wrong"}], "must be an array"),
        ([{"title": "T", "bullets": [5]}], "must be a string"),
        ([{"title": "T", "image": "/etc/passwd"}], "accepts only title and bullets"),
        ([{"title": "T", "bullets": ["emoji \U0001F642"]}], "unsupported by the bundled font"),
        ([{"title": "T", "bullets": ["bad\x00text"]}], "unsupported control"),
    ],
)
def test_invalid_or_unbounded_specs_are_rejected(spec, match):
    with pytest.raises(renderers.RenderError, match=match):
        renderers.render_slide_deck("", spec, "pdf")


def test_deck_title_and_total_text_are_bounded(monkeypatch):
    with pytest.raises(renderers.RenderError, match="deck title is too long"):
        renderers.render_slide_deck("x" * 256, SAMPLE, "pptx")
    import rich_slides
    monkeypatch.setattr(rich_slides, "MAX_DECK_CHARS", 3)
    with pytest.raises(renderers.RenderError, match="deck text is too long"):
        renderers.render_slide_deck("", SAMPLE, "pdf")


def test_maximum_slide_count_is_supported():
    spec = [{"title": f"Slide {n}"} for n in range(slides.MAX_SLIDES)]
    assert len(PdfReader(io.BytesIO(renderers.render_slide_deck("", spec, "pdf"))).pages) == slides.MAX_SLIDES
