"""Round-trip the workbook features and reject formulas outside the subset."""

from __future__ import annotations

import io
import zipfile

import pytest
from openpyxl import Workbook, load_workbook

from renderers import RenderError, _cell, render_xlsx
from workbook_features import WorkbookFeatureError, apply_workbook_features


def _sheet(spec):
    wb = Workbook()
    ws = wb.active
    ws.title = "Revenue"
    if spec.get("columns"):
        ws.append([_cell(value) for value in spec["columns"]])
    for row in spec.get("rows", []):
        ws.append([_cell(value) for value in row])
    return wb, ws


def _reopen(wb):
    buffer = io.BytesIO()
    wb.save(buffer)
    return load_workbook(io.BytesIO(buffer.getvalue())), buffer.getvalue()


def test_explicit_formula_round_trips_while_other_cells_keep_native_types():
    spec = {
        "columns": ["Month", "Revenue", "Active", "Zip", "Source"],
        "rows": [["Jan", 100, True, "01234", "=SUM(B2:B3)"], ["Feb", 125.5, False, "00123", "-offset"], ["Total", "", None, "", ""]],
        "formulas": [{"cell": "b4", "expression": " sum($b$2:b3) "}],
    }
    wb, ws = _sheet(spec)
    apply_workbook_features(ws, spec)
    reopened, _ = _reopen(wb)
    result = reopened.active
    assert result["B4"].value == "=SUM($B$2:B3)"
    assert result["B4"].data_type == "f"
    assert result["B2"].value == 100 and result["B2"].data_type == "n"
    assert result["B3"].value == 125.5 and result["B3"].data_type == "n"
    assert result["C2"].value is True and result["C2"].data_type == "b"
    assert result["D2"].value == "01234" and result["D2"].data_type == "s"
    assert result["E2"].value == "'=SUM(B2:B3)" and result["E2"].data_type == "s"
    assert result["E3"].value == "'-offset" and result["E3"].data_type == "s"
    assert result.auto_filter.ref == "A1:E4"


def test_renderer_integration_recalculates_explicit_formulas_and_preserves_chart():
    data = render_xlsx([{
        "name": "Revenue",
        "columns": ["Month", "Revenue"],
        "rows": [["Jan", 100], ["Feb", 125], ["Total", ""]],
        "formulas": [{"cell": "B4", "expression": "SUM(B2:B3)"}],
        "chart": {"title": "Revenue", "kind": "bar", "category_column": 1, "value_column": 2},
    }])
    wb = load_workbook(io.BytesIO(data))
    assert wb["Revenue"]["B4"].value == "=SUM(B2:B3)"
    assert len(wb["Revenue"]._charts) == 1
    assert wb.calculation.fullCalcOnLoad is True
    assert wb.calculation.forceFullCalc is True
    assert wb.calculation.calcMode == "auto"


@pytest.mark.parametrize("expression", [
    "=ROUND(AVERAGE(B2:B3),2)", "SUM(B2:B3)+MIN(B2:B3)-MAX(B2:B3)",
    "ABS(-B2)/COUNT(B2:B3)", "(B2+B3)*.5^2+10%", "1e3+2.5E-2", "+B2", "1+2+3+4",
])
def test_numeric_formula_subset_survives_save_and_reopen(expression):
    spec = {"columns": ["Month", "Revenue"], "rows": [["Jan", 100], ["Feb", 125], ["Total", ""]], "formulas": [{"cell": "B4", "expression": expression}]}
    wb, ws = _sheet(spec)
    apply_workbook_features(ws, spec)
    reopened, _ = _reopen(wb)
    assert reopened.active["B4"].data_type == "f"
    assert reopened.active["B4"].value.startswith("=")


@pytest.mark.parametrize("kind, expected_tag", [("bar", b"barChart"), ("line", b"lineChart")])
def test_chart_contains_actual_data_references_and_reopens(kind, expected_tag):
    spec = {"columns": ["Month", "Revenue"], "rows": [["Jan", 100], ["Feb", 125]], "chart": {"title": "Revenue", "kind": kind, "category_column": 1, "value_column": 2}}
    wb, ws = _sheet(spec)
    apply_workbook_features(ws, spec)
    reopened, data = _reopen(wb)
    charts = reopened.active._charts
    assert len(charts) == 1
    assert charts[0].series[0].val.numRef.f == "'Revenue'!$B$2:$B$3"
    assert charts[0].series[0].cat.numRef.f == "'Revenue'!$A$2:$A$3"
    assert charts[0].series[0].tx.strRef.f == "'Revenue'!B1"
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        xml = archive.read("xl/charts/chart1.xml")
        assert expected_tag in xml and b"Revenue" in xml


@pytest.mark.parametrize("expression", [
    'HYPERLINK("https://evil.example","click")', "WEBSERVICE(A2)", "INDIRECT(A2)",
    "OFFSET(A2,0,0)", "CALL(A2)", "REGISTER.ID(A2)", "EXEC(A2)", "OLE(A2)",
    "[evil.xlsx]Sheet1!A1", "'Sheet 2'!A1", "Sheet1!A1", "cmd|' /C calc'!A1",
    "https://evil.example", "SUM(A:A)", "SUM(1:3)", "SUM(A2#)", "@A2", "A2&A3",
    "SUM()", "SUM(A2,)", "ROUND(A2)", "ABS(A2,A3)", "SUM(A2", "A2 A3",
    "1+", "=", "SUM(A4:A2)", "A1000000", "XFD2", "SUM(A2:B500000)",
])
def test_unsafe_or_malformed_formulas_are_refused(expression):
    spec = {"columns": ["Month", "Revenue"], "rows": [["Jan", 100], ["Feb", 125], ["Total", ""]], "formulas": [{"cell": "B4", "expression": expression}]}
    _, ws = _sheet(spec)
    with pytest.raises(WorkbookFeatureError):
        apply_workbook_features(ws, spec)
    assert ws["B4"].data_type != "f"


@pytest.mark.parametrize("target", ["B5", "C2", "A0", "A1:B2", "Sheet1!A1", "[file]A1", 2, None])
def test_formula_target_cannot_enlarge_or_escape_grid(target):
    spec = {"columns": ["Month", "Revenue"], "rows": [["Jan", 100]], "formulas": [{"cell": target, "expression": "SUM(B2)"}]}
    _, ws = _sheet(spec)
    with pytest.raises(WorkbookFeatureError):
        apply_workbook_features(ws, spec)
    assert (ws.max_row, ws.max_column) == (2, 2)


def test_all_features_are_validated_before_any_cell_is_changed():
    spec = {"columns": ["Month", "Revenue"], "rows": [["Jan", 100], ["Total", ""]], "formulas": [{"cell": "B3", "expression": "SUM(B2)"}, {"cell": "A3", "expression": "INDIRECT(A2)"}]}
    _, ws = _sheet(spec)
    with pytest.raises(WorkbookFeatureError):
        apply_workbook_features(ws, spec)
    assert ws["B3"].value == "" and ws["A3"].value == "Total"
    assert ws.auto_filter.ref is None


def test_duplicate_formula_targets_are_refused_instead_of_overwritten():
    spec = {"columns": ["Value"], "rows": [[1], [""]], "formulas": [{"cell": "A3", "expression": "SUM(A2)"}, {"cell": "$a$3", "expression": "42"}]}
    _, ws = _sheet(spec)
    with pytest.raises(WorkbookFeatureError, match="duplicated"):
        apply_workbook_features(ws, spec)
    assert ws["A3"].value == ""


@pytest.mark.parametrize("expression", ["1+" * 5000 + "1", "(" * 40 + "1" + ")" * 40, "+" * 40 + "1", "1^(" * 40 + "1" + ")" * 40])
def test_oversized_or_deep_formulas_are_rejected_without_truncation(expression):
    spec = {"columns": ["Value"], "rows": [[""]], "formulas": [{"cell": "A2", "expression": expression}]}
    _, ws = _sheet(spec)
    with pytest.raises(WorkbookFeatureError):
        apply_workbook_features(ws, spec)
    assert ws["A2"].value == ""


@pytest.mark.parametrize("patch", [
    {"kind": "scatter"}, {"category_column": 0}, {"value_column": 3},
    {"category_column": True}, {"value_column": 2.0}, {"category_column": 2},
    {"title": "x" * 201},
])
def test_bad_chart_is_refused_before_formula_mutation(patch):
    chart = {"title": "Revenue", "kind": "bar", "category_column": 1, "value_column": 2, **patch}
    spec = {"columns": ["Month", "Revenue"], "rows": [["Jan", 100], ["Total", ""]], "chart": chart, "formulas": [{"cell": "B3", "expression": "SUM(B2)"}]}
    _, ws = _sheet(spec)
    with pytest.raises(WorkbookFeatureError):
        apply_workbook_features(ws, spec)
    assert ws["B3"].value == "" and not ws._charts


@pytest.mark.parametrize("columns, rows", [(["Month", "Revenue"], []), ([], [["Jan", 100], ["Feb", 125]])])
def test_chart_needs_explicit_headers_and_data(columns, rows):
    spec = {"columns": columns, "rows": rows, "chart": {"kind": "line", "category_column": 1, "value_column": 2}}
    _, ws = _sheet(spec)
    with pytest.raises(WorkbookFeatureError):
        apply_workbook_features(ws, spec)


@pytest.mark.parametrize("value", ["100", "=SUM(B2)", "text", True, False, float("nan"), float("inf"), float("-inf")])
def test_chart_refuses_non_numeric_value_cells_before_any_mutation(value):
    spec = {"columns": ["Month", "Revenue"], "rows": [["Jan", value], ["Total", ""]], "formulas": [{"cell": "B3", "expression": "100"}], "chart": {"kind": "bar", "category_column": 1, "value_column": 2}}
    _, ws = _sheet(spec)
    with pytest.raises(WorkbookFeatureError, match="chart values"):
        apply_workbook_features(ws, spec)
    assert ws["B3"].value == "" and not ws._charts


def test_chart_allows_numbers_and_explicit_formulas_with_optional_blanks():
    spec = {"columns": ["Month", "Revenue"], "rows": [["Jan", 0], ["Feb", None], ["Mar", ""], ["Apr", "placeholder"]], "formulas": [{"cell": "B5", "expression": "42"}], "chart": {"kind": "line", "category_column": 1, "value_column": 2}}
    wb, ws = _sheet(spec)
    apply_workbook_features(ws, spec)
    reopened, _ = _reopen(wb)
    assert reopened.active["B5"].value == "=42"
    assert len(reopened.active._charts) == 1


def test_chart_needs_some_numeric_or_formula_data():
    spec = {"columns": ["Month", "Revenue"], "rows": [["Jan", None], ["Feb", ""]], "chart": {"kind": "bar", "category_column": 1, "value_column": 2}}
    _, ws = _sheet(spec)
    with pytest.raises(WorkbookFeatureError, match="at least one numeric"):
        apply_workbook_features(ws, spec)


def test_renderer_maps_bad_chart_values_to_render_error():
    with pytest.raises(RenderError, match="chart values"):
        render_xlsx([{"columns": ["Month", "Revenue"], "rows": [["Jan", True]], "chart": {"kind": "bar", "category_column": 1, "value_column": 2}}])


@pytest.mark.parametrize("formulas", [
    [{"cell": "B3", "expression": "B3+1"}],
    [{"cell": "B3", "expression": "SUM(B2:B3)"}],
    [{"cell": "A3", "expression": "B3"}, {"cell": "B3", "expression": "A3"}],
    [{"cell": "A3", "expression": "SUM(B2:B3)"}, {"cell": "B3", "expression": "A3+1"}],
])
def test_circular_formula_dependencies_are_refused_before_mutation(formulas):
    spec = {"columns": ["Month", "Revenue"], "rows": [["Jan", 100], ["Total", ""]], "formulas": formulas}
    _, ws = _sheet(spec)
    with pytest.raises(WorkbookFeatureError, match="circular"):
        apply_workbook_features(ws, spec)
    assert ws["A3"].value == "Total" and ws["B3"].value == ""


def test_long_acyclic_formula_chain_uses_no_recursive_graph_traversal():
    formulas = [{"cell": f"A{row}", "expression": f"A{row + 1}+1"} for row in range(2, 1001)]
    formulas.append({"cell": "A1001", "expression": "1"})
    spec = {"columns": ["Value"], "rows": [[""] for _ in range(1000)], "formulas": formulas}
    wb, ws = _sheet(spec)
    apply_workbook_features(ws, spec)
    reopened, _ = _reopen(wb)
    assert reopened.active["A2"].value == "=A3+1"
    assert reopened.active["A1001"].value == "=1"


def test_too_many_formulas_are_refused_without_dropping_any():
    spec = {"columns": ["Value"], "rows": [[""] for _ in range(1001)], "formulas": [{"cell": f"A{row}", "expression": "1"} for row in range(2, 1003)]}
    _, ws = _sheet(spec)
    with pytest.raises(WorkbookFeatureError, match="at most 1000"):
        apply_workbook_features(ws, spec)
    assert ws["A2"].value == "" and ws["A1002"].value == ""


@pytest.mark.parametrize("formulas", ["A2=1", {}, ["A2=1"], [{"cell": "A2", "expression": 42}]])
def test_invalid_formula_shapes_fail_closed(formulas):
    spec = {"columns": ["Value"], "rows": [[""]], "formulas": formulas}
    _, ws = _sheet(spec)
    with pytest.raises(WorkbookFeatureError):
        apply_workbook_features(ws, spec)


def test_empty_sheet_does_not_get_a_filter_or_an_invented_data_grid():
    spec = {"columns": [], "rows": []}
    wb, ws = _sheet(spec)
    apply_workbook_features(ws, spec)
    reopened, _ = _reopen(wb)
    assert reopened.active.auto_filter.ref is None
    assert reopened.active["A1"].value is None


def test_cached_values_and_cross_sheet_formula_dependencies_reopen_without_excel():
    data = render_xlsx([
        {"name": "Summary", "columns": ["Value", "Rounded"], "rows": [["", ""], ["", ""]], "formulas": [
            {"cell": "A2", "expression": "SUM('Sales Q1'!B2:B3)"},
            {"cell": "B2", "expression": "ROUND(A2/3,2)"},
            {"cell": "A3", "expression": "1/0"}, {"cell": "B3", "expression": "ABS(A3)"}]},
        {"name": "Sales Q1", "columns": ["Month", "Amount"], "rows": [["Jan", 10], ["Feb", 20]]},
    ])
    cached = load_workbook(io.BytesIO(data), data_only=True)
    assert cached["Summary"]["A2"].value == 30
    assert cached["Summary"]["B2"].value == 10
    assert cached["Summary"]["A3"].value == "#DIV/0!"
    assert cached["Summary"]["B3"].value == "#DIV/0!"
    assert load_workbook(io.BytesIO(data))["Summary"]["A2"].value == "=SUM('Sales Q1'!B2:B3)"


@pytest.mark.parametrize("expression, expected", [
    ("2^3^2", 64), ("2^-2^3", .015625),
    ("2^(3^2)", 512), ("(2^3)^2", 64), ("-2^2", 4),
])
def test_formula_cache_uses_excel_power_associativity(expression, expected):
    data = render_xlsx([{"columns": ["Value"], "rows": [[""]],
        "formulas": [{"cell": "A2", "expression": expression}]}])
    assert load_workbook(io.BytesIO(data), data_only=True).active["A2"].value == expected
    assert load_workbook(io.BytesIO(data)).active["A2"].value == "=" + expression


@pytest.mark.parametrize("expression, expected", [
    ("COUNT(A2:A4)", 1), ("COUNT(A3)", 0), ("COUNT(A3,5)", 1),
])
def test_count_cache_ignores_errors_and_non_numeric_cells(expression, expected):
    data = render_xlsx([{"columns": ["Values", "Count"], "rows": [[10, ""], ["", ""], [True, ""]],
        "formulas": [{"cell": "A3", "expression": "1/0"}, {"cell": "B2", "expression": expression}]}])
    cached = load_workbook(io.BytesIO(data), data_only=True).active
    assert cached["A3"].value == "#DIV/0!"
    assert cached["B2"].value == expected


def test_cross_sheet_cycles_are_refused_and_no_workbook_bytes_escape():
    with pytest.raises(RenderError, match="circular"):
        render_xlsx([
            {"name": "One", "columns": ["Value"], "rows": [[""]], "formulas": [{"cell": "A2", "expression": "Two!A2"}]},
            {"name": "Two", "columns": ["Value"], "rows": [[""]], "formulas": [{"cell": "A2", "expression": "One!A2"}]},
        ])


@pytest.mark.parametrize("kind, columns", [("bar", [2, 3]), ("line", [2, 3]), ("pie", [2])])
def test_multiseries_and_pie_charts_include_numeric_caches(kind, columns):
    data = render_xlsx([{"name": "Data", "columns": ["Month", "Sales", "Cost"], "rows": [["Jan", 10, 8], ["Feb", "", 9]],
        "formulas": [{"cell": "B3", "expression": "B2*2"}],
        "chart": {"kind": kind, "category_column": 1, "value_columns": columns}}])
    chart = load_workbook(io.BytesIO(data)).active._charts[0]
    assert len(chart.series) == len(columns)
    assert [point.v for point in chart.series[0].val.numRef.numCache.pt] == [10, 20]


def test_formats_preserve_native_types_and_iso_dates():
    import datetime
    data = render_xlsx([{"columns": ["Amount", "Share", "Date"], "rows": [[10, .125, "2026-10-08"]], "formats": [
        {"range": "A2", "kind": "currency", "currency": "EUR", "precision": 2},
        {"range": "B2", "kind": "percent", "precision": 1}, {"range": "C2", "kind": "date"}]}])
    ws = load_workbook(io.BytesIO(data)).active
    assert ws["A2"].value == 10 and ws["A2"].number_format == '"EUR" #,##0.00'
    assert ws["B2"].value == .125 and ws["B2"].number_format == "0.0%"
    assert ws["C2"].value == datetime.datetime(2026, 10, 8) and ws["C2"].number_format == "yyyy-mm-dd"


@pytest.mark.parametrize("format_spec", [
    {"range": "A2:B1", "kind": "number"}, {"range": "A3", "kind": "number"},
    {"range": "A2", "kind": "currency", "currency": "BOGUS"},
    {"range": "A2", "kind": "number", "precision": True},
    {"range": "A2", "kind": "date"},
])
def test_malformed_or_unsupported_formats_fail_before_render(format_spec):
    with pytest.raises(RenderError):
        render_xlsx([{"columns": ["Value", "Other"], "rows": [["not a date", 2]], "formats": [format_spec]}])


def _formula_result(expression, rows=None):
    data = render_xlsx([{"name": "Data", "columns": ["Value", "Amount", "Result"],
        "rows": rows or [[10, 2, ""], [20, 4, ""], [30, 6, ""]],
        "formulas": [{"cell": "C2", "expression": expression}]}])
    return load_workbook(io.BytesIO(data), data_only=True)["Data"]["C2"], data


@pytest.mark.parametrize("expression, expected", [
    ('SUM(TRUE,"2")', 3), ('COUNT(TRUE,"2")', 2),
    ('MIN(5,TRUE,"2")', 1), ('MAX(1,FALSE,"2")', 2),
    ('AVERAGE(TRUE,"2")', 1.5), ('MIN(FALSE,2)', 0),
    ('AVERAGE(FALSE,2)', 1), ('COUNT(1/0,"bad",TRUE)', 1),
    ('SUM("bad",2)', "#VALUE!"), ('AVERAGE("bad",2)', "#VALUE!"),
    ('MIN("bad",2)', "#VALUE!"), ('MAX("bad",2)', "#VALUE!"),
])
def test_aggregates_coerce_direct_scalar_arguments_like_excel(expression, expected):
    # COUNT/MIN/MAX distinguish direct literals from referenced cells. SUM's
    # documented SUM("5",15,TRUE)=21 example follows the same rule. These
    # cached results were also checked in Excel using an unsaved workbook.
    result, data = _formula_result(expression)
    assert result.value == expected
    assert result.data_type == ("e" if isinstance(expected, str) else "n")
    assert load_workbook(io.BytesIO(data))["Data"]["C2"].value == "=" + expression


@pytest.mark.parametrize("expression, expected", [
    ('SUM(A2:A4)', 5), ('COUNT(A2:A4)', 1),
    ('SUM(IF(TRUE,A2,0))', 0), ('SUM(IF(TRUE,A3,0))', 0),
    ('COUNT(IF(TRUE,A3,0))', 0), ('SUM(IF(FALSE,A2,"2"))', 2),
    ('COUNT(IF(TRUE,"2",0))', 1),
    ('SUM(IF(TRUE,IF(FALSE,0,A3),0))', 0),
    ('SUM(IF(TRUE,A2,0)+0)', 1),
])
def test_aggregates_preserve_selected_if_reference_semantics(expression, expected):
    # Selecting a cell through IF preserves its reference status. Applying
    # arithmetic to that selection instead produces a direct scalar value.
    assert _formula_result(expression, [[True, 2, ""], ["2", 4, ""], [5, 6, ""]])[0].value == expected


@pytest.mark.parametrize("expression, expected", [
    ('IF(TRUE,A2,1)', 0), ('A2', 0),
    ('SUM(IF(TRUE,A2,0))', 0), ('COUNT(IF(TRUE,A2,0))', 0),
])
def test_selected_empty_cells_cache_zero_without_becoming_counted_arguments(expression, expected):
    assert _formula_result(expression, [[None, 2, ""], [5, 4, ""]])[0].value == expected


@pytest.mark.parametrize("expression, expected, data_type", [
    ('IF(A2>=10,"Met Target","Below Target")', "Met Target", "s"),
    ('IF(A2<10,"Met Target","Below Target")', "Below Target", "s"),
    ('IF(TRUE,"a,b(1) = WEBSERVICE(""literal"")","unused")', 'a,b(1) = WEBSERVICE("literal")', "s"),
    ('IF(FALSE,1/0,42)', 42, "n"),
    ('IF(TRUE,1/0,"safe")', "#DIV/0!", "e"),
    ('IF(TRUE,FALSE,1/0)', False, "b"),
    ('IF(FALSE,1/0,TRUE)', True, "b"),
    ('A2+B2>=12', True, "b"),
    ('A2<>10', False, "b"),
    ('"Quarter One"="quarter one"', True, "b"),
    ('"1"=1', False, "b"),
    ('IF("invalid condition",1,2)', "#VALUE!", "e"),
    ('IF("FALSE",1,IF(A2=10,7,8))', 7, "n"),
])
def test_business_formulas_round_trip_typed_caches_and_preserve_literals(expression, expected, data_type):
    result, data = _formula_result(expression)
    assert result.value == expected and result.data_type == data_type
    reopened = load_workbook(io.BytesIO(data))
    assert reopened["Data"]["C2"].value == "=" + expression
    assert reopened["Data"]["C2"].data_type == "f"
    assert reopened.calculation.fullCalcOnLoad is True
    assert reopened.calculation.forceFullCalc is True
    assert reopened.calculation.calcMode == "auto"


def test_empty_string_cache_keeps_standard_string_type_and_formula():
    # openpyxl treats an empty cached <v> as None. Keep the standard OOXML
    # string cache so Excel and consumers still distinguish it from zero.
    import xml.etree.ElementTree as ET
    result, data = _formula_result('IF(TRUE,"","unused")')
    assert result.value is None and result.data_type == "str"
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        root = ET.fromstring(archive.read("xl/worksheets/sheet1.xml"))
    ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
    cell = root.find(f"{ns}sheetData/{ns}row/{ns}c[@r='C2']")
    assert cell.attrib["t"] == "str" and cell.find(f"{ns}v") is not None
    assert cell.find(f"{ns}f").text == 'IF(TRUE,"","unused")'


@pytest.mark.parametrize("expression, expected", [
    ("1+2*3=7", True), ("1+2*3<>7", False), ("2^3<9", True),
    ("2^3<=8", True), ("2^3>9", False), ("2^3>=8", True),
    ("TRUE=1", False), ('"abc"<"ABD"', True),
    ('IF(TRUE,A2:A4,0)', "#VALUE!"), ('IF(FALSE,1/0,1/0)', "#DIV/0!"),
])
def test_comparison_precedence_boolean_types_and_selected_if_errors(expression, expected):
    assert _formula_result(expression)[0].value == expected


def test_string_cache_escaping_and_error_looking_literals_remain_strings():
    result, _ = _formula_result('IF(TRUE," <a> & ""Quoted"" ",1)')
    assert result.value == ' <a> & "Quoted" ' and result.data_type == "s"
    result, _ = _formula_result('IF(TRUE,"#DIV/0!",1)')
    assert result.value == "#DIV/0!" and result.data_type == "s"


def test_native_error_cells_propagate_through_comparisons():
    result, _ = _formula_result("IF(A2=1,7,8)", [["#DIV/0!", 1, ""]])
    assert result.value == "#DIV/0!" and result.data_type == "e"


@pytest.mark.parametrize("expression, expected", [
    ('COUNTIF(A2:A4,">=20")', 2), ('COUNTIF(A2:A4,20)', 1),
    ('COUNTIF(A2:A4,A3)', 1), ('COUNTIF(A2:A4,"<>20")', 2),
    ('SUMIF(A2:A4,">10",B2:B4)', 10), ('SUMIF(A2:A4,">10")', 50),
    ('COUNTIF(A2,10)', 1), ('SUMIF(A2,10,B2)', 2),
    ('SUMIF(A2:A4,IF(A2=10,">=20","=10"),B2:B4)', 10),
])
def test_conditional_counts_and_sums_use_scalar_or_range_criteria(expression, expected):
    assert _formula_result(expression)[0].value == expected


def test_criteria_numeric_strings_and_boolean_cells_preserve_underlying_types():
    data = render_xlsx([{"columns": ["Value", "Result"], "rows": [[1, ""], ["1", ""], [True, ""], [False, ""]],
        "formulas": [{"cell": "B2", "expression": "COUNTIF(A2:A5,1)"},
            {"cell": "B3", "expression": "COUNTIF(A2:A5,TRUE)"}]}])
    ws = load_workbook(io.BytesIO(data), data_only=True).active
    assert ws["B2"].value == 2 and ws["B3"].value == 1
    assert ws["A2"].data_type == "n" and ws["A3"].data_type == "s"
    assert ws["A4"].value is True and ws["A4"].data_type == "b"


def test_cross_sheet_text_criteria_and_formula_results_feed_summary():
    data = render_xlsx([
        {"name": "Summary", "columns": ["Region", "Sales", "Matches"], "rows": [["north", "", ""]], "formulas": [
            {"cell": "B2", "expression": "SUMIF('Sales Q1'!A2:A4,A2,'Sales Q1'!B2:B4)"},
            {"cell": "C2", "expression": "COUNTIF('Sales Q1'!A2:A4,A2)"}]},
        {"name": "Sales Q1", "columns": ["Region", "Sales"], "rows": [["NORTH", 10], ["South", 20], ["North", ""]],
            "formulas": [{"cell": "B4", "expression": "IF(B2=10,B2*3,1/0)"}]},
    ])
    ws = load_workbook(io.BytesIO(data), data_only=True)["Summary"]
    assert ws["B2"].value == 40 and ws["C2"].value == 2


@pytest.mark.parametrize("criteria, expected", [
    ('"N*"', 2), ('"n?rth"', 2), ('"~*"', 1), ('"Q~?"', 1),
    ('"A~~B"', 1), ('""', 1), ('"<>"', 5),
])
def test_case_insensitive_wildcards_escaped_metacharacters_and_blank_criteria(criteria, expected):
    rows = [[text, 1, ""] for text in ["NORTH", "North", "*", "Q?", "A~B", None]]
    result, _ = _formula_result(f"COUNTIF(A2:A7,{criteria})", rows)
    assert result.value == expected


def test_sumif_propagates_only_matched_sum_errors():
    data = render_xlsx([{"columns": ["Region", "Amount", "Summary"], "rows": [["North", 10, ""], ["South", "", ""]],
        "formulas": [{"cell": "B3", "expression": "1/0"},
            {"cell": "C2", "expression": 'SUMIF(A2:A3,"North",B2:B3)'},
            {"cell": "C3", "expression": 'SUMIF(A2:A3,"South",B2:B3)'}]}])
    ws = load_workbook(io.BytesIO(data), data_only=True).active
    assert ws["C2"].value == 10 and ws["C3"].value == "#DIV/0!"


@pytest.mark.parametrize("expression", [
    "IF(TRUE,1)", "IF(TRUE,1,2,3)", "COUNTIF(A2:A3)", "COUNTIF(1,2)",
    "SUMIF(A2:A3,1,B2)", "SUMIF(A2:B2,1,A2:A3)", "SUMIF(A2:A3,1,2)",
    'IF(FALSE,WEBSERVICE("https://evil.invalid"),1)',
    'IF(FALSE,INDIRECT("A2"),1)', 'IF(TRUE,"unterminated,1)',
    'IF(TRUE,"a""",REGISTER.ID(A2))', 'IF(TRUE,"\x00",1)',
    'IF(TRUE,"\ud800",1)', 'IF(TRUE,"\uffff",1)',
    'IF(TRUE,1,\'[external.xlsx]Data\'!A2)', 'IF(TRUE,1,cmd|\'/c calc\'!A2)',
])
def test_new_grammar_refuses_shape_errors_and_unsafe_dead_branches_before_mutation(expression):
    spec = {"columns": ["Value", "Amount", "Result"], "rows": [[1, 2, ""], [3, 4, ""]],
        "formulas": [{"cell": "C2", "expression": expression}]}
    _, ws = _sheet(spec)
    with pytest.raises(WorkbookFeatureError): apply_workbook_features(ws, spec)
    assert ws["C2"].value == "" and ws["C2"].data_type != "f"


def test_static_cycle_detection_includes_unselected_if_branches():
    with pytest.raises(RenderError, match="circular"):
        _formula_result("IF(TRUE,1,C2)")


def test_if_does_not_read_or_evaluate_an_unselected_range(monkeypatch):
    import workbook_calculation as calculation
    monkeypatch.setattr(calculation, "MAX_CELL_READS", 0)
    assert _formula_result("IF(TRUE,7,SUM(A2:A4))")[0].value == 7
    with pytest.raises(RenderError, match="cell reads"):
        _formula_result("IF(FALSE,7,SUM(A2:A4))")


def test_criteria_length_and_error_criteria_are_cached_as_errors():
    assert _formula_result('COUNTIF(A2:A4,"' + "a" * 256 + '")')[0].value == "#VALUE!"
    assert _formula_result("SUMIF(A2:A4,1/0,B2:B4)")[0].value == "#DIV/0!"


def test_formula_evaluation_has_independent_work_and_elapsed_time_limits(monkeypatch):
    import workbook_calculation as calculation
    monkeypatch.setattr(calculation, "MAX_EVALUATION_STEPS", 5)
    with pytest.raises(RenderError, match="work limit"):
        _formula_result("1+2+3+4")
    monkeypatch.setattr(calculation, "MAX_EVALUATION_STEPS", 1_000_000)
    monkeypatch.setattr(calculation, "MAX_CALCULATION_SECONDS", -1)
    with pytest.raises(RenderError, match="time limit"):
        _formula_result("1")


def test_dependency_graph_has_its_own_work_limit_before_mutation(monkeypatch):
    import workbook_calculation as calculation
    monkeypatch.setattr(calculation, "MAX_DEPENDENCY_CHECKS", 1)
    spec = {"columns": ["Value", "Result"], "rows": [[10, ""], [20, ""]],
        "formulas": [{"cell": "B2", "expression": "SUM(A2:A3)"}]}
    _, ws = _sheet(spec)
    with pytest.raises(WorkbookFeatureError, match="dependencies.*work limit"):
        apply_workbook_features(ws, spec)
    assert ws["B2"].value == ""


def test_pathological_wildcards_consume_work_budget_without_regex_backtracking(monkeypatch):
    import workbook_calculation as calculation
    monkeypatch.setattr(calculation, "MAX_EVALUATION_STEPS", 500)
    with pytest.raises(RenderError, match="work limit"):
        _formula_result('COUNTIF(A2,"*' + "a" * 100 + 'b")', [["a" * 10000, 1, ""]])


def test_native_named_table_round_trips_with_chart_cached_formulas_and_filter():
    import xml.etree.ElementTree as ET
    data = render_xlsx([{"name": "Sales", "table_name": "QuarterlySales", "columns": ["Month", "Sales"],
        "rows": [["Jan", 10], ["Feb", ""]], "formulas": [{"cell": "B3", "expression": "IF(B2>0,B2*2,0)"}],
        "chart": {"kind": "bar", "category_column": 1, "value_column": 2}}])
    ws = load_workbook(io.BytesIO(data)).active
    table = ws.tables["QuarterlySales"]
    assert table.ref == "A1:B3" and table.autoFilter.ref == "A1:B3"
    assert [column.name for column in table.tableColumns] == ["Month", "Sales"]
    assert table.tableStyleInfo.name == "TableStyleMedium2" and table.tableStyleInfo.showRowStripes is True
    assert len(ws._charts) == 1 and ws["B3"].value == "=IF(B2>0,B2*2,0)"
    assert load_workbook(io.BytesIO(data), data_only=True).active["B3"].value == 20
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        assert "xl/tables/table1.xml" in archive.namelist()
        table_xml = ET.fromstring(archive.read("xl/tables/table1.xml"))
        assert table_xml.attrib["displayName"] == "QuarterlySales" and table_xml.attrib["ref"] == "A1:B3"
        assert b"tablePart" in archive.read("xl/worksheets/sheet1.xml")
        assert b"/table" in archive.read("xl/worksheets/_rels/sheet1.xml.rels")


def test_tables_are_explicit_and_ordinary_workbook_filter_behavior_is_preserved():
    data = render_xlsx([{"columns": ["Month", "Sales"], "rows": [["Jan", 10]]}])
    ws = load_workbook(io.BytesIO(data)).active
    assert not ws.tables and ws.auto_filter.ref == "A1:B2"
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        assert not any(path.startswith("xl/tables/") for path in archive.namelist())


@pytest.mark.parametrize("name", [None, 1, "", " ", "1Sales", "Quarterly Sales", "Sales/Rows", "a" * 256,
    "A1", "XFD1048576", "R", "c", "RC", "R1C1", "r20c12", "R1C1Suffix", "R3C", "RC4", "_xlnm.Print_Area"])
def test_invalid_or_reserved_table_names_fail_before_formulas_are_written(name):
    spec = {"table_name": name, "columns": ["Month", "Sales"], "rows": [["Jan", ""]],
        "formulas": [{"cell": "B2", "expression": "10"}]}
    _, ws = _sheet(spec)
    with pytest.raises(WorkbookFeatureError, match="table_name"):
        apply_workbook_features(ws, spec)
    assert ws["B2"].value == "" and not ws.tables and ws.auto_filter.ref is None


@pytest.mark.parametrize("columns, rows", [
    ([], [["Jan", 1]]), (["Month", "Sales"], []), (["Month"], [["Jan", 1]]),
    (["Month", ""], [["Jan", 1]]), (["Month", " "], [["Jan", 1]]),
    (["Sales", "sales"], [[1, 2]]), (["Sales", "Sales "], [[1, 2]]),
    (["a" * 256], [[1]]), ([123], [[1]]), ([True], [[1]]), (["#DIV/0!"], [[1]]),
    (["\ud800"], [[1]]), (["\uffff"], [[1]]),
])
def test_tables_require_complete_literal_unique_bounded_headers(columns, rows):
    with pytest.raises(RenderError, match="headers"):
        render_xlsx([{"table_name": "SalesTable", "columns": columns, "rows": rows}])


def test_formula_targets_and_date_conversions_cannot_change_table_headers():
    with pytest.raises(RenderError, match="headers cannot be formula"):
        render_xlsx([{"table_name": "SalesTable", "columns": ["Sales"], "rows": [[1]],
            "formulas": [{"cell": "A1", "expression": '"Revenue"'}]}])
    with pytest.raises(RenderError, match="headers must remain literal"):
        render_xlsx([{"table_name": "SalesTable", "columns": ["2026-10-08"], "rows": [[1]],
            "formats": [{"range": "A1", "kind": "date"}]}])


def test_table_names_are_unique_case_insensitively_across_sheets():
    with pytest.raises(RenderError, match="unique across"):
        render_xlsx([{"name": name, "table_name": table_name, "columns": ["Value"], "rows": [[1]]}
            for name, table_name in [("One", "SalesTable"), ("Two", "salestable")]])


def test_table_names_do_not_conflict_with_existing_defined_names():
    from openpyxl.workbook.defined_name import DefinedName
    spec = {"table_name": "SalesTable", "columns": ["Value"], "rows": [[1]]}
    wb, ws = _sheet(spec)
    wb.defined_names.add(DefinedName("salestable", attr_text="'Revenue'!$A$2"))
    with pytest.raises(WorkbookFeatureError, match="unique across"):
        apply_workbook_features(ws, spec)
    assert not ws.tables


def test_valid_unicode_and_maximum_names_headers_and_literal_leaders_round_trip():
    data = render_xlsx([{"name": "Sales", "table_name": "Ventes_Été", "columns": ["=Label", "+Amount"], "rows": [["Jan", 1]]},
        {"name": "Bounds", "table_name": "_" + "a" * 254, "columns": ["b" * 255], "rows": [[1]]}])
    wb = load_workbook(io.BytesIO(data))
    assert [column.name for column in wb["Sales"].tables["Ventes_Été"].tableColumns] == ["'=Label", "'+Amount"]
    assert wb["Sales"]["A1"].data_type == "s" and wb["Sales"]["B1"].data_type == "s"
    assert len(wb["Bounds"].tables["_" + "a" * 254].tableColumns[0].name) == 255


def test_native_tables_keep_existing_grid_row_limits(monkeypatch):
    import renderers
    monkeypatch.setattr(renderers, "MAX_ROWS_PER_SHEET", 2)
    with pytest.raises(RenderError, match="too many rows"):
        render_xlsx([{"table_name": "SalesTable", "columns": ["Value"], "rows": [[1], [2], [3]]}])


def test_multiple_distinct_tables_preserve_native_parts_and_global_ids():
    data = render_xlsx([{"name": name, "table_name": table_name, "columns": ["Value"], "rows": [[1]]}
        for name, table_name in [("One", "FirstTable"), ("Two", "SecondTable")]])
    wb = load_workbook(io.BytesIO(data))
    assert wb["One"].tables["FirstTable"].id == 1
    assert wb["Two"].tables["SecondTable"].id == 2


def test_authenticated_xlsx_route_preserves_named_table_formulas_styles_and_chart(client, auth):
    response = client.post("/render", headers=auth, json={"format": "xlsx", "sheets": [{
        "name": "Sales", "table_name": "SalesTable", "columns": ["Month", "Amount"],
        "rows": [["October", 100], ["November", None]],
        "formulas": [{"cell": "B3", "expression": "IF(B2>=100,B2*1.5,0)"}],
        "formats": [{"range": "B2:B3", "kind": "currency", "currency": "USD", "precision": 2}],
        "chart": {"kind": "bar", "category_column": 1, "value_column": 2},
    }]})
    assert response.status_code == 200
    workbook = load_workbook(io.BytesIO(response.content))
    sheet = workbook["Sales"]
    assert sheet.tables["SalesTable"].ref == "A1:B3"
    assert sheet.tables["SalesTable"].autoFilter.ref == "A1:B3"
    assert sheet["B3"].value == "=IF(B2>=100,B2*1.5,0)"
    assert sheet["B3"].number_format == '"USD" #,##0.00' and len(sheet._charts) == 1
    assert load_workbook(io.BytesIO(response.content), data_only=True)["Sales"]["B3"].value == 150


def test_authenticated_xlsx_route_omitted_table_name_keeps_ordinary_sheet(client, auth):
    response = client.post("/render", headers=auth, json={"format": "xlsx", "sheets": [{
        "columns": ["Value"], "rows": [[1]]}]})
    assert response.status_code == 200
    sheet = load_workbook(io.BytesIO(response.content)).active
    assert not sheet.tables and sheet.auto_filter.ref == "A1:A2"


@pytest.mark.parametrize("patch", [
    {"table_name": None}, {"table_name": 123}, {"table_name": "A1"}, {"table_name": "R1C1"},
    {"columns": ["Value", "value"]}, {"columns": ["Value", ""]},
])
def test_authenticated_xlsx_route_refuses_invalid_native_table_specs_without_file_bytes(client, auth, patch):
    response = client.post("/render", headers=auth, json={"format": "xlsx", "sheets": [{
        "table_name": "SalesTable", "columns": ["Value", "Other"], "rows": [[1, 2]], **patch}]})
    assert response.status_code in (400, 422)
    assert response.headers["content-type"].startswith("application/json")
    assert not response.content.startswith(b"PK")
