"""Validated workbook formulas, cached calculation, native charts and formats."""
from __future__ import annotations
import datetime
import io
import math
import re
import zipfile
import xml.etree.ElementTree as ET
from typing import Any
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.chart.data_source import NumData, NumVal
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from workbook_calculation import CELL, Parser, coordinate, formula_order, calculate, CellError, FormulaError as WorkbookFeatureError

MAX_FORMULAS = 1000

def _dimensions(ws, spec):
    columns, rows = spec.get("columns") or [], spec.get("rows") or []
    return len(rows) + bool(columns), max([len(columns)] + [len(row) for row in rows])

def _chart(spec, ws, dimensions, rows, formula_targets):
    if not isinstance(spec, dict): raise WorkbookFeatureError("sheet chart must be an object")
    kind = spec.get("kind")
    if kind not in ("bar", "line", "pie"): raise WorkbookFeatureError("chart kind must be bar, line or pie")
    single, multiple = spec.get("value_column"), spec.get("value_columns")
    if (single is None) == (multiple is None): raise WorkbookFeatureError("chart needs exactly one value_column or value_columns")
    values = [single] if single is not None else multiple
    if not isinstance(values, list) or not 1 <= len(values) <= 8 or any(type(c) is not int for c in values) or len(set(values)) != len(values): raise WorkbookFeatureError("chart needs 1-8 distinct value columns")
    columns = [spec.get("category_column"), *values]
    if any(type(column) is not int or not 1 <= column <= dimensions[1] for column in columns): raise WorkbookFeatureError("chart columns must be integers inside the existing sheet grid")
    if columns[0] in values: raise WorkbookFeatureError("chart category and value columns must be different")
    if kind == "pie" and len(values) != 1: raise WorkbookFeatureError("pie chart requires exactly one value series")
    if dimensions[0] < 2: raise WorkbookFeatureError("chart needs a header row and at least one data row")
    for column in values:
        has_value = False
        for index, row in enumerate(rows, start=2):
            if f"{get_column_letter(column)}{index}" in formula_targets: has_value = True; continue
            value = row[column - 1] if len(row) >= column else None
            if value is None or value == "": continue
            if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value): raise WorkbookFeatureError("chart values must be finite numbers, explicit formulas or blanks")
            if kind == "pie" and value < 0: raise WorkbookFeatureError("pie values must be non-negative")
            has_value = True
        if not has_value: raise WorkbookFeatureError("chart needs at least one numeric value or explicit formula")
    title = spec.get("title", "")
    if not isinstance(title, str) or len(title) > 200: raise WorkbookFeatureError("chart title must be a string of at most 200 characters")
    chart = BarChart() if kind == "bar" else LineChart() if kind == "line" else PieChart()
    if title: chart.title = title
    for column in values: chart.add_data(Reference(ws, min_col=column, min_row=1, max_row=dimensions[0]), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=columns[0], min_row=2, max_row=dimensions[0]))
    chart._droplet_value_columns = values
    return chart

def _formats(ws, specs, dimensions):
    if not isinstance(specs, list) or len(specs) > 64: raise WorkbookFeatureError("sheet supports at most 64 formats")
    changes, seen = [], set()
    for spec in specs:
        if not isinstance(spec, dict): raise WorkbookFeatureError("each format must be an object")
        bounds = spec.get("range", "").split(":") if isinstance(spec.get("range"), str) else []
        if len(bounds) not in (1, 2): raise WorkbookFeatureError("format range must be local A1 or A1:B2")
        start, end = coordinate(bounds[0], dimensions), coordinate(bounds[-1], dimensions)
        a, b = CELL.fullmatch(start), CELL.fullmatch(end)
        left, top, right, bottom = column_index_from_string(a[1]), int(a[2]), column_index_from_string(b[1]), int(b[2])
        if right < left or bottom < top: raise WorkbookFeatureError("format range must run top-left to bottom-right")
        precision = spec.get("precision", 2)
        if type(precision) is not int or not 0 <= precision <= 8: raise WorkbookFeatureError("format precision must be an integer from 0 to 8")
        kind = spec.get("kind")
        number = "#,##0" + ("." + "0" * precision if precision else "")
        if kind == "number": code = number
        elif kind == "percent": code = "0" + ("." + "0" * precision if precision else "") + "%"
        elif kind == "date": code = "yyyy-mm-dd"
        elif kind == "currency":
            currency = spec.get("currency", "USD")
            if currency not in ("USD", "EUR", "GBP", "JPY", "CAD", "AUD"): raise WorkbookFeatureError("unsupported currency format")
            code = f'"{currency}" ' + number
        else: raise WorkbookFeatureError("format kind must be number, currency, percent or date")
        for row in range(top, bottom + 1):
            for column in range(left, right + 1):
                key = row, column
                if key in seen: raise WorkbookFeatureError("format ranges must not overlap")
                seen.add(key)
                if len(seen) > 100_000: raise WorkbookFeatureError("format ranges exceed 100000 cells")
                value = ws.cell(row, column).value
                if kind == "date" and isinstance(value, str) and value:
                    try:
                        value = datetime.datetime.fromisoformat(value) if "T" in value else datetime.date.fromisoformat(value)
                        if isinstance(value, datetime.datetime) and value.tzinfo: raise ValueError("timezone")
                    except ValueError: raise WorkbookFeatureError("date format strings must be ISO dates without timezone")
                changes.append((row, column, code, value))
    return changes

def _table(ws, spec, dimensions, formula_targets, formats):
    if "table_name" not in spec: return None
    name = spec["table_name"]
    if not isinstance(name, str) or not 1 <= len(name) <= 255 or not (name[0].isalpha() or name[0] == "_") or any(not (char.isalpha() or char.isdecimal() or char in "_.") for char in name):
        raise WorkbookFeatureError("table_name must be 1-255 letters, digits, underscores or periods, starting with a letter or underscore")
    folded = name.casefold()
    if CELL.fullmatch(name) or folded in ("r", "c", "rc") or re.match(r"R[1-9][0-9]*C[1-9][0-9]*", name, re.IGNORECASE) or re.fullmatch(r"R[1-9][0-9]*C|RC[1-9][0-9]*", name, re.IGNORECASE) or folded.startswith(("_xlnm.", "_xlpm.", "_xlfn.")):
        raise WorkbookFeatureError("table_name cannot be a cell reference, R1C1 reference or reserved Excel name")
    existing = {table_name.casefold() for sheet in ws.parent for table_name in sheet.tables}
    existing.update(defined_name.casefold() for defined_name in ws.parent.defined_names)
    if folded in existing: raise WorkbookFeatureError("table names must be unique across the workbook")
    columns = spec.get("columns") or []
    if not columns or not spec.get("rows") or len(columns) != dimensions[1]:
        raise WorkbookFeatureError("named tables require explicit headers for every column and at least one data row")
    if any(CELL.fullmatch(target)[2] == "1" for target in formula_targets):
        raise WorkbookFeatureError("named table headers cannot be formula targets")
    if any(row == 1 and not isinstance(value, str) for row, _, _, value in formats):
        raise WorkbookFeatureError("named table headers must remain literal text after formatting")
    headers = [ws.cell(1, column).value for column in range(1, dimensions[1] + 1)]
    if any(not isinstance(raw, str) or not isinstance(header, str) or not header.strip() or len(header) > 255 or ws.cell(1, column).data_type != "s" for column, (raw, header) in enumerate(zip(columns, headers), 1)):
        raise WorkbookFeatureError("named table headers must be nonblank literal strings of at most 255 characters")
    if any(not (char in "\t\n\r" or 0x20 <= ord(char) <= 0xD7FF or 0xE000 <= ord(char) <= 0xFFFD or 0x10000 <= ord(char) <= 0x10FFFF) for header in headers for char in header):
        raise WorkbookFeatureError("named table headers must contain valid XML text")
    if len({header.strip().casefold() for header in headers}) != len(headers):
        raise WorkbookFeatureError("named table headers must be unique ignoring case and surrounding spaces")
    table = Table(displayName=name, ref=f"A1:{get_column_letter(dimensions[1])}{dimensions[0]}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
    table._initialise_columns()
    for column, header in zip(table.tableColumns, headers): column.name = header
    return table

def apply_workbook_features(ws: Any, spec: dict[str, Any]) -> None:
    dimensions = _dimensions(ws, spec)
    workbook = ws.parent
    all_dimensions = {sheet.title: getattr(sheet, "_droplet_dimensions", (sheet.max_row, sheet.max_column)) for sheet in workbook}
    all_dimensions[ws.title] = dimensions
    raw = spec.get("formulas")
    if raw is None: raw = []
    if not isinstance(raw, list) or len(raw) > MAX_FORMULAS: raise WorkbookFeatureError("sheet supports at most 1000 explicit formulas")
    validated, targets = {}, set()
    for formula in raw:
        if not isinstance(formula, dict): raise WorkbookFeatureError("each formula must be an object")
        target = coordinate(formula.get("cell"), dimensions)
        if target in targets: raise WorkbookFeatureError("formula targets must not be duplicated")
        targets.add(target)
        match = CELL.fullmatch(target)
        validated[(ws.title, column_index_from_string(match[1]), int(match[2]))] = Parser(formula.get("expression"), ws.title, all_dimensions).parse()
    formula_order(validated)
    chart = None
    if spec.get("chart") is not None:
        if not spec.get("columns"): raise WorkbookFeatureError("chart needs explicit column headers")
        chart = _chart(spec["chart"], ws, dimensions, spec.get("rows") or [], targets)
    formats = _formats(ws, spec.get("formats") or [], dimensions)
    table = _table(ws, spec, dimensions, targets, formats)
    for row, column, number_format, value in formats:
        ws.cell(row, column).number_format = number_format
        ws.cell(row, column).value = value
    for (_, column, row), (expression, _, _) in validated.items(): ws.cell(row, column).value = expression
    ws._droplet_formulas = validated
    ws._droplet_dimensions = dimensions
    if table is not None: ws.add_table(table)
    elif spec.get("columns") and spec.get("rows") and dimensions[1]: ws.auto_filter.ref = f"A1:{get_column_letter(dimensions[1])}{dimensions[0]}"
    if chart is not None: ws.add_chart(chart, f"{get_column_letter(dimensions[1] + 2)}2")

def workbook_caches(workbook):
    formulas = {target: formula for ws in workbook for target, formula in getattr(ws, "_droplet_formulas", {}).items()}
    if len(formulas) > 1000: raise WorkbookFeatureError("workbook supports at most 1000 explicit formulas in total")
    caches = calculate(workbook, formulas)
    for ws in workbook:
        for chart in ws._charts:
            for series, column in zip(chart.series, chart._droplet_value_columns):
                points = []
                for index, row in enumerate(range(2, ws._droplet_dimensions[0] + 1)):
                    value = caches.get((ws.title, column, row), ws.cell(row, column).value)
                    if value in (None, "") or isinstance(value, CellError): continue
                    if isinstance(value, bool) or not isinstance(value, (float, int)) or not math.isfinite(value): raise WorkbookFeatureError("calculated chart series must contain numbers or blanks")
                    if isinstance(chart, PieChart) and value < 0: raise WorkbookFeatureError("calculated pie values must be non-negative")
                    points.append(NumVal(idx=index, v=value))
                if not points: raise WorkbookFeatureError("calculated chart series has no numeric values")
                series.val.numRef.numCache = NumData(ptCount=ws._droplet_dimensions[0] - 1, pt=points)
    return caches

def inject_formula_caches(data: bytes, workbook, caches) -> bytes:
    """Keep formulas plus portable OOXML caches for read-only clients."""
    namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    ET.register_namespace("", namespace)
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as source, zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as target:
        sheet_paths = {f"xl/worksheets/sheet{index}.xml": ws.title for index, ws in enumerate(workbook, 1)}
        for info in source.infolist():
            contents = source.read(info.filename)
            if info.filename in sheet_paths:
                sheet = sheet_paths[info.filename]
                root = ET.fromstring(contents)
                for cell in root.findall(f".//{{{namespace}}}c"):
                    match = CELL.fullmatch(cell.attrib.get("r", ""))
                    if not match: continue
                    key = sheet, column_index_from_string(match[1]), int(match[2])
                    if key not in caches: continue
                    value = caches[key]
                    if isinstance(value, CellError): cell.set("t", "e"); text = value.code
                    elif isinstance(value, bool): cell.set("t", "b"); text = "1" if value else "0"
                    elif value is None: cell.set("t", "n"); text = "0"
                    elif isinstance(value, str): cell.set("t", "str"); text = value
                    else: cell.set("t", "n"); text = str(value)
                    cached = cell.find(f"{{{namespace}}}v")
                    if cached is None: cached = ET.SubElement(cell, f"{{{namespace}}}v")
                    cached.text = text
                contents = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            target.writestr(info, contents)
    return output.getvalue()
