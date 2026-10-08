"""One-shot, credential-free data analysis host for the existing sandbox.

Trusted loaders decode bounded CSV/XLSX bytes before user Python starts. User
code receives JSON tables, the routine runner's stdlib allowlist, and narrowly
scoped artifact helpers. Artifacts are returned as bytes, never written into a
workspace. Linux Landlock/seccomp seals this child before source decoding or
user code; the container is an additional boundary.
"""
from __future__ import annotations

import base64
import csv
import datetime
import html
import io
import json
import math
import re
import sys
import zipfile
from pathlib import Path

# -I removes the script directory; only trusted sibling code is imported here.
sys.path.insert(0, str(Path(__file__).resolve().parent))
import runner
from analysis_isolation import seal_analysis

# Load the approved stdlib while this trusted host still has its ordinary
# importer. Their private dependencies vary by Python minor version; user code
# receives only the guarded importer, after the kernel boundary is installed.
for _module_name in runner.ALLOWED_MODULES:
    __import__(_module_name)

# The installed trusted loader initializes Python's MIME registry on import,
# which may consult /etc/mime.types. Finish that dependency initialization
# before sealing. No user bytes/code are processed until AFTER the kernel
# boundary; never grant the analysis child read access to /etc to satisfy it.
import openpyxl
import encodings.cp437  # ZIP filename codec; no lazy module stat after sealing.
import encodings.utf_8_sig  # CSV BOM decoding; same trusted initialization.

MAX_SOURCE_BYTES = 3 * 1024 * 1024
MAX_CELLS = 100_000
MAX_ROWS = 10_000
MAX_COLUMNS = 200
MAX_ARTIFACT_BYTES = 512_000
MAX_ARTIFACTS = 8
MAX_LOG_BYTES = 8192
NAME = re.compile(r"^[A-Za-z0-9][A-Za-z0-9_. -]{0,79}$")


def _cell(value):
    if isinstance(value, (datetime.datetime, datetime.date, datetime.time)):
        return value.isoformat()
    if value is None or isinstance(value, (str, bool, int)):
        return value
    if isinstance(value, float) and math.isfinite(value):
        return value
    raise ValueError("source contains an unsupported or non-finite cell value")


def _sheet(name, values):
    rows = []
    count = 0
    for row in values:
        if len(rows) >= MAX_ROWS + 1 or len(row) > MAX_COLUMNS:
            raise ValueError(f"{name}: maximum {MAX_ROWS} data rows and {MAX_COLUMNS} columns")
        count += len(row)
        if count > MAX_CELLS:
            raise ValueError(f"{name}: more than {MAX_CELLS} cells")
        rows.append([_cell(v) for v in row])
    if not rows:
        return {"name": name, "columns": [], "rows": []}
    # Preserve duplicate/blank headers and row positions: never quietly drop data.
    columns = ["" if v is None else str(v) for v in rows[0]]
    return {"name": name, "columns": columns, "rows": rows[1:]}


def _load(source):
    name = source.get("name", "")
    if not isinstance(name, str) or len(name) > 255:
        raise ValueError("invalid source name")
    raw = base64.b64decode(source.get("contentBase64", ""), validate=True)
    if len(raw) > MAX_SOURCE_BYTES:
        raise ValueError("source exceeds 3 MiB")
    fmt = source.get("format")
    if fmt == "csv":
        text = raw.decode("utf-8-sig", errors="strict")
        # Explicit comma-separated CSV. Do not guess delimiters or numeric types.
        sheet = _sheet("CSV", csv.reader(io.StringIO(text), strict=True))
        return {"name": name, "sheets": [sheet]}, []
    if fmt != "xlsx":
        raise ValueError("sources must be UTF-8 CSV or XLSX")
    # Reject expansion bombs before the spreadsheet library allocates XML trees.
    with zipfile.ZipFile(io.BytesIO(raw)) as archive:
        entries = archive.infolist()
        if len(entries) > 1000 or sum(e.file_size for e in entries) > 16 * 1024 * 1024:
            raise ValueError("XLSX uncompressed content exceeds 16 MiB / 1000 entries")
        if any(e.file_size > 8 * 1024 * 1024 or e.flag_bits & 1 for e in entries):
            raise ValueError("XLSX has an oversized or encrypted entry")
    workbook = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True, keep_links=False)
    try:
        if len(workbook.worksheets) > 10:
            raise ValueError("XLSX has more than 10 sheets")
        sheets = []
        total = 0
        for worksheet in workbook.worksheets:
            if (worksheet.max_row or 0) > MAX_ROWS + 1 or (worksheet.max_column or 0) > MAX_COLUMNS:
                raise ValueError(f"{worksheet.title}: maximum {MAX_ROWS} data rows and {MAX_COLUMNS} columns")
            sheet = _sheet(worksheet.title, worksheet.iter_rows(values_only=True))
            total += sum(len(r) for r in sheet["rows"]) + len(sheet["columns"])
            if total > MAX_CELLS:
                raise ValueError(f"XLSX has more than {MAX_CELLS} cells in total")
            sheets.append(sheet)
        return {"name": name, "sheets": sheets}, [
            f"{name}: formula cells use saved cached values; absent caches are null. Formulas are not recalculated."
        ]
    finally:
        workbook.close()


class Artifacts:
    def __init__(self):
        self.items = []

    def _add(self, name, extension, mime, contents):
        if not isinstance(name, str) or not NAME.fullmatch(name) or ".." in name or not name.lower().endswith(extension):
            raise ValueError(f"artifact name must be a simple filename ending in {extension}")
        if len(self.items) >= MAX_ARTIFACTS:
            raise ValueError(f"maximum {MAX_ARTIFACTS} artifacts")
        if any(item["name"] == name for item in self.items):
            raise ValueError("duplicate artifact name")
        if len(contents) > MAX_ARTIFACT_BYTES:
            raise ValueError(f"artifact exceeds {MAX_ARTIFACT_BYTES} bytes")
        self.items.append({"name": name, "mimeType": mime, "contentBase64": base64.b64encode(contents).decode("ascii")})
        return name

    def csv(self, name, columns, rows):
        if not isinstance(columns, (list, tuple)) or not isinstance(rows, (list, tuple)):
            raise ValueError("emit_csv requires columns and rows lists")
        if len(columns) > MAX_COLUMNS or len(rows) > MAX_ROWS or len(columns) * (len(rows) + 1) > MAX_CELLS:
            raise ValueError("CSV artifact exceeds table limits")
        text = io.StringIO()
        writer = csv.writer(text)

        def safe(value):
            value = _cell(value)
            # Excel must treat generated strings as data, not formulas.
            if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@", "\t", "\r")):
                return "'" + value
            return value

        writer.writerow([safe(v) for v in columns])
        for row in rows:
            if not isinstance(row, (list, tuple)) or len(row) != len(columns):
                raise ValueError("CSV rows must match columns")
            writer.writerow([safe(v) for v in row])
            if text.tell() > MAX_ARTIFACT_BYTES:
                raise ValueError("CSV artifact exceeds byte limit")
        return self._add(name, ".csv", "text/csv", text.getvalue().encode("utf-8"))

    def chart(self, name, title, labels, values, kind="bar"):
        if kind not in ("bar", "line") or not isinstance(labels, (list, tuple)) or not isinstance(values, (list, tuple)):
            raise ValueError("emit_chart supports bar or line with labels/values lists")
        if not 1 <= len(labels) <= 100 or len(labels) != len(values):
            raise ValueError("chart requires 1-100 matching labels and values")
        if not isinstance(title, str) or len(title) > 160 or any(not isinstance(v, str) or len(v) > 80 for v in labels):
            raise ValueError("chart title/labels must be bounded strings")
        if any(ord(char) < 32 for text in [title, *labels] for char in text):
            raise ValueError("chart title/labels must not contain control characters")
        if any(isinstance(v, bool) or not isinstance(v, (int, float)) or not math.isfinite(v) for v in values):
            raise ValueError("chart values must be finite numbers")
        low, high = min(0, min(values)), max(0, max(values))
        if low == high:
            high = low + 1
        if not math.isfinite(high - low):
            raise ValueError("chart value range is too large")
        left, top, width, height = 90, 70, 740, 350
        step = width / len(values)
        y = lambda value: top + height * (high - value) / (high - low)
        zero = y(0)
        pieces = [
            '<svg xmlns="http://www.w3.org/2000/svg" width="900" height="520" viewBox="0 0 900 520">',
            '<rect width="900" height="520" fill="white"/>',
            f'<text x="90" y="38" font-family="sans-serif" font-size="22">{html.escape(title)}</text>',
            f'<path d="M {left} {top} V {top + height} M {left} {zero} H {left + width}" stroke="#64748b" fill="none"/>',
        ]
        for tick in range(5):
            value = low + (high - low) * tick / 4
            pieces.append(f'<text x="80" y="{y(value) + 5:.2f}" text-anchor="end" font-family="sans-serif" font-size="12">{value:.4g}</text>')
        points = []
        for i, value in enumerate(values):
            x = left + step * (i + 0.5)
            points.append(f"{x:.2f},{y(value):.2f}")
            if kind == "bar":
                pieces.append(f'<rect x="{x - step * .35:.2f}" y="{min(y(value), zero):.2f}" width="{step * .7:.2f}" height="{abs(y(value) - zero):.2f}" fill="#4f46e5"/>')
            # All values are charted; reduce only axis label density explicitly.
            if i % max(1, math.ceil(len(values) / 12)) == 0:
                pieces.append(f'<text x="{x:.2f}" y="445" transform="rotate(35 {x:.2f} 445)" font-family="sans-serif" font-size="11">{html.escape(labels[i])}</text>')
        if kind == "line":
            pieces.append(f'<polyline points="{" ".join(points)}" fill="none" stroke="#4f46e5" stroke-width="3"/>')
        pieces.append('</svg>')
        return self._add(name, ".svg", "image/svg+xml", "".join(pieces).encode("utf-8"))


class Log:
    def __init__(self):
        self.text = ""
        self.truncated = False

    def print(self, *values, sep=" ", end="\n"):
        value = sep.join(str(v) for v in values) + end
        remaining = MAX_LOG_BYTES - len(self.text.encode("utf-8"))
        raw = value.encode("utf-8")
        self.text += raw[:max(0, remaining)].decode("utf-8", "ignore")
        self.truncated = self.truncated or len(raw) > remaining


def analyze(request):
    try:
        seal_analysis(request.get("scratchDir", ""))
        runner._apply_limits(int(request.get("maxMemoryBytes") or 256 * 1024 * 1024))
        tables, warnings = [], []
        sources = request.get("sources") or []
        if not isinstance(sources, list) or len(sources) > 4:
            raise ValueError("maximum 4 sources")
        for source in sources:
            table, source_warnings = _load(source)
            tables.append(table)
            warnings.extend(source_warnings)
        artifacts, log = Artifacts(), Log()
        result = runner._run(request, {
            "tables": tables, "emit_csv": artifacts.csv, "emit_chart": artifacts.chart, "print": log.print,
        })
        if "error" in result:
            return result
        # Unlike routine transforms, reject NaN/Infinity: JSON must be portable.
        json.dumps(result["output"], allow_nan=False)
        return {**result, "stdout": log.text, "stdoutTruncated": log.truncated, "artifacts": artifacts.items,
                "sources": [{"name": t["name"], "sheets": [{"name": s["name"], "columns": s["columns"], "rowCount": len(s["rows"])} for s in t["sheets"]]} for t in tables],
                "warnings": warnings}
    except MemoryError:
        return {"error": "analysis exceeded its memory budget"}
    except Exception as exc:
        return {"error": f"{type(exc).__name__}: {exc}"}


if __name__ == "__main__":
    try:
        result = analyze(json.load(sys.stdin))
    except Exception as exc:
        result = {"error": f"invalid analysis request: {exc}"}
    sys.stdout.write(json.dumps(result, separators=(",", ":"), allow_nan=False))
    sys.stdout.flush()
